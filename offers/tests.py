"""Angebote: python manage.py test offers --settings=tabele.settings_test"""
import io
import json
import shutil
import tempfile
from datetime import datetime, timezone as dt_tz
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.utils import timezone

from crm.models import Customer
from email_assistant.models import EmailAccount, EmailSettings, EmailThread
from email_assistant.sync import save_message
from inventory.models import Product
from offers import pricing, services
from offers.models import Offer, OfferLine, OfferRevision
from sales.models import SalesOrder, SalesOrderLine

ME = "sales@sevskiy.de"
D = Decimal


def _all_text(path) -> str:
    """Увесь текст DOCX, включно з вкладеними таблицями бланка."""
    import re
    import zipfile
    return " ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", zipfile.ZipFile(path).read("word/document.xml").decode()))


class OffersBase(TestCase):
    def setUp(self):
        from config.models import SystemSettings
        self.media = tempfile.mkdtemp()
        self._media = override_settings(MEDIA_ROOT=self.media)
        self._media.enable()
        s = SystemSettings.objects.get_or_create(pk=1)[0]
        s.is_onboarding_complete = True
        s.default_currency = "EUR"
        s.save()
        self.user = User.objects.create_superuser("slava", "s@x.y", "p")
        self.client.force_login(self.user)
        self.customer = Customer.objects.create(name="Alex Buyer", company="Client GmbH",
                                                email="buyer@client.example", addr_street="Hauptstr. 1",
                                                addr_zip="10115", addr_city="Berlin", country="DE")
        self.ant = Product.objects.create(sku="ANT-100", name="Antenna 2.4 GHz", sale_price=D("12.50"),
                                          purchase_price=D("6.00"))
        self.cab = Product.objects.create(sku="CAB-7", name="RF Cable 0.5 m", sale_price=D("3.10"))

    def tearDown(self):
        self._media.disable()
        shutil.rmtree(self.media, ignore_errors=True)


class PricingTests(OffersBase):
    def test_sources_in_order(self):
        from shop.models import Shop, ShopListing, ShopPriceTier, ShopSettings
        shop = Shop.objects.create(name="Website", slug="webshop", is_default=True, currency="EUR")
        listing = ShopListing.objects.create(shop=shop, product=self.ant, price=D("10.00"))
        ShopPriceTier.objects.create(listing=listing, min_qty=100, unit_price=D("8.00"))
        q = pricing.quote_price(self.ant, 150, "EUR")
        self.assertEqual((q.unit_price, q.source), (D("8.0000"), OfferLine.SHOP))
        self.assertIn("від 100 шт", q.note)

        q = pricing.quote_price(self.cab, 5, "EUR")                     # немає позиції магазину
        self.assertEqual((q.unit_price, q.source), (D("3.1000"), OfferLine.SALE_PRICE))

        chip = Product.objects.create(sku="IC-55", name="Chip", base_price_currency="USD",
                                      base_prices=[{"min_qty": 1, "unit_price": "2.00"},
                                                   {"min_qty": 500, "unit_price": "1.50"}])
        q = pricing.quote_price(chip, 600, "EUR")
        self.assertEqual(q.source, OfferLine.MANUAL)                       # курсу USD ще немає
        self.assertIn("немає курсу USD→EUR", q.note)
        ss = ShopSettings.get()
        ss.fx_rates = {"USD": "0.92"}
        ss.save()
        q = pricing.quote_price(chip, 600, "EUR")
        self.assertEqual((q.unit_price, q.source), (D("1.3800"), OfferLine.BASE))
        self.assertIn("від 500 шт", q.note)
        nothing = Product.objects.create(sku="X-1", name="No price")
        self.assertIsNone(pricing.quote_price(nothing, 1, "EUR").unit_price)
        self.assertIsNone(pricing.quote_price(self.cab, 1, "GBP").unit_price)   # курсу немає — без вигадок
        q = pricing.quote_price(self.cab, 1, "USD")                            # є курс USD — перерахунок
        self.assertEqual((q.unit_price, q.source), (D("3.3700"), OfferLine.SALE_PRICE))

    def test_totals_are_exact(self):
        offer = services.new_offer(customer=self.customer, user=self.user, shipping_cost=D("15.00"))
        services.add_line(offer, product=self.ant, qty=3, unit_price=D("12.345"), discount_percent=D("10"))
        services.add_line(offer, sku="SERVICE", description="Kalibrierung", qty=1, unit_price=D("99.99"))
        offer.refresh_from_db()
        # 3 × 12.345 = 37.035 − 10 % = 33.3315 → 33.33; + 99.99 + 15.00 = 148.32; ПДВ 19 % = 28.18
        self.assertEqual([l.line_total for l in offer.lines.all()], [D("33.33"), D("99.99")])
        self.assertEqual((offer.net_total, offer.vat_total, offer.gross_total), (D("148.32"), D("28.18"), D("176.50")))


class LifecycleTests(OffersBase):
    def test_number_document_lock_and_new_version(self):
        year = timezone.localdate().year
        offer = services.new_offer(customer=self.customer, user=self.user)
        self.assertEqual(offer.number, f"A-{year}-001")
        self.assertEqual(services.new_offer(user=self.user).number, f"A-{year}-002")
        self.assertIn("Hauptstr. 1", offer.billing_address)
        services.add_line(offer, product=self.ant, qty=10)
        rev = services.generate_document(offer, self.user)
        self.assertTrue(rev.docx_file)
        self.assertEqual(rev.snapshot["lines"][0]["unit_price"], "12.5000")
        services.mark_sent(rev, "<m1@x>")
        offer.refresh_from_db()
        self.assertTrue(offer.is_locked)
        with self.assertRaises(services.OfferError):
            services.add_line(offer, product=self.cab)
        self.ant.sale_price = D("99")                                         # каталог змінився після відправки
        self.ant.save()
        self.assertEqual(services.generate_document(offer).snapshot["lines"][0]["unit_price"], "12.5000")
        services.new_version(offer, self.user)
        self.assertEqual((offer.revision, offer.status), (2, Offer.DRAFT))
        services.add_line(offer, product=self.cab, qty=2)
        rev2 = services.generate_document(offer, self.user)
        self.assertEqual((rev2.revision, len(rev2.snapshot["lines"])), (2, 2))
        rev.refresh_from_db()
        self.assertEqual((len(rev.snapshot["lines"]), rev.sent_at is not None), (1, True))
        self.assertEqual(rev2.file_name, f"Angebot_{offer.number}-V2")

    def test_document_content(self):
        from docx import Document
        offer = services.new_offer(customer=self.customer, user=self.user, shipping_cost=D("1000"))
        services.add_line(offer, product=self.ant, qty=100, unit_price=D("12.5"), lead_time="3 Wochen")
        services.add_line(offer, sku="X", description="Sonderanfertigung", qty=1)          # без ціни
        rev = services.generate_document(offer)
        text = "\n".join(p.text for p in Document(rev.docx_file.path).paragraphs)
        cells = "\n".join(c.text for t in Document(rev.docx_file.path).tables for r in t.rows for c in r.cells)
        self.assertIn(f"Angebotsnr.: \t{offer.number}", text)              # бланк Commercial Invoice як Quotation
        self.assertIn("Lieferadresse (Empfänger/Käufer):", cells)
        self.assertIn("Absender (Verkäufer):", _all_text(rev.docx_file.path))   # блок продавця — вкладена таблиця
        self.assertIn("Rechnungsadresse: Client GmbH, Hauptstr. 1, 10115 Berlin, Germany", text)
        self.assertIn("Summe netto", cells)                                   # рядок таблиці: 1 250 + доставка 1 000
        self.assertIn("2.250,00", cells)
        self.assertIn("Verpackung und Versand", cells)
        self.assertIn("Gesamtbetrag: 2.677,50 EUR", text)                    # ПДВ 19 % — під таблицею
        self.assertIn("Antenna 2.4 GHz ANT-100", cells)
        self.assertIn("auf Anfrage", cells)
        self.assertIn("3 Wochen", cells)
        self.assertIn("Alle Preise verstehen sich zzgl. MwSt.", text)
        self.assertNotIn("Zahlungsinformationen", text)                       # реквізитів не задано — блоку немає
        self.assertNotIn("{{", text + cells)
        from offers.models import OfferSettings
        s = OfferSettings.get()
        s.payment_info = "Beneficiary: Sevskiy GmbH\nIBAN: DE00 0000"
        s.save()
        rev = services.generate_document(offer)
        text = "\n".join(p.text for p in Document(rev.docx_file.path).paragraphs)
        self.assertIn("Zahlungsinformationen:", text)
        self.assertIn("Beneficiary: Sevskiy GmbH\nIBAN: DE00 0000", text)

    def test_accept_creates_order(self):
        offer = services.new_offer(customer=self.customer, user=self.user)
        services.add_line(offer, product=self.ant, qty=4, unit_price=D("10"), discount_percent=D("5"))
        services.add_line(offer, sku="SVC", description="Service", qty=1, unit_price=D("50"))
        with self.assertRaises(services.OfferError):
            services.accept(offer)                                            # ще не надіслано
        services.mark_sent(services.generate_document(offer))
        order, skipped = services.accept(Offer.objects.get(pk=offer.pk), self.user)
        self.assertEqual((order.source, order.order_number, skipped), ("angebot", offer.number, [2]))
        line = SalesOrderLine.objects.get(order=order)
        self.assertEqual((line.qty, line.total_price, line.unit_price, line.currency),
                         (D("4.000"), D("38.00"), D("9.5000"), "EUR"))
        offer.refresh_from_db()
        self.assertEqual((offer.status, offer.order), (Offer.ACCEPTED, order))
        self.assertEqual(services.accept(offer)[0], order)                   # повтор — без другого замовлення
        import os
        self.assertTrue(os.path.exists(os.path.join(self.media, "orders", "angebot", offer.number,
                                                    f"Angebot_{offer.number}.docx")))

    def test_invalid_status_transition(self):
        offer = services.new_offer(user=self.user)
        with self.assertRaises(services.OfferError):
            services.set_status(offer, Offer.REJECTED)
        services.set_status(offer, Offer.REVIEW)
        services.set_status(offer, Offer.APPROVED)
        self.assertEqual(offer.status, Offer.APPROVED)


class FromMailAndQuoteTests(OffersBase):
    def setUp(self):
        super().setUp()
        self.account = EmailAccount.objects.create(
            user=self.user, email_address=ME, imap_username=ME, imap_password="x",
            smtp_username=ME, smtp_password="x", imap_folder_inbox="INBOX")
        es = EmailSettings.get_for_user(self.user)
        es.deadline_detection = es.telegram_notify_new = es.mark_read_on_server = False
        es.save()

    def mail(self, subject, body, frm="buyer@client.example"):
        return save_message(self.account, {
            "uid": 1, "message_id": "<q1@client.example>", "in_reply_to": "", "references": [],
            "subject": subject, "from_email": frm, "from_name": "Alex Buyer", "to_emails": [ME], "cc_emails": [],
            "body_text": body, "body_html": "", "attachments": [], "is_read": False,
            "sent_at": datetime(2026, 10, 1, 9, 0, tzinfo=dt_tz.utc)}, "INBOX")

    def test_detect_items(self):
        found = services.detect_items("Please quote 500 pcs ANT-100\nCAB-7 x 20\nalso IC-99999 (unknown)\nant-100 again")
        self.assertEqual([(p.sku, q) for p, q in found], [("ANT-100", D("500")), ("CAB-7", D("20"))])
        self.assertEqual([(p.sku, q) for p, q in services.detect_items("Bitte Angebot für CAB-7")], [("CAB-7", None)])

    def test_from_thread_button(self):
        from projects import services as project_services
        m = self.mail("Anfrage Antennen", "Hello,\nplease quote 500 pcs ANT-100 and CAB-7 x 20.\nThanks")
        project = project_services.create_project(name="Antennen", user=self.user, threads=[m.thread])
        r = self.client.post(f"/offers/api/from-thread/{m.thread_id}/")
        d = r.json()
        self.assertTrue(d["ok"])
        offer = Offer.objects.get(pk=d["id"])
        self.assertEqual((offer.customer, offer.project, offer.source_message), (self.customer, project, m))
        self.assertEqual([(l.sku, l.qty, l.unit_price) for l in offer.lines.all()],
                         [("ANT-100", D("500.000"), D("12.5000")), ("CAB-7", D("20.000"), D("3.1000"))])
        self.assertEqual(offer.contact_email, "buyer@client.example")
        pv = self.client.get(f"/email/thread/{m.thread_id}/preview/")
        self.assertContains(pv, "pvOfferFromThread")
        other = User.objects.create_user("bob", "b@x.y", "p", is_staff=True)
        self.client.force_login(other)
        self.assertEqual(self.client.post(f"/offers/api/from-thread/{m.thread_id}/").status_code, 404)

    def test_from_quote_action(self):
        quote = SalesOrder.objects.create(source="webshop", order_number="Q-77", document_type="QUOTE",
                                          email="buyer@client.example", customer_key=self.customer.external_key,
                                          affects_stock=False)
        SalesOrderLine.objects.create(order=quote, product=self.ant, qty=D("50"), unit_price=D("11"), currency="EUR")
        r = self.client.post("/admin/sales/salesorder/", {"action": "create_offer_from_quote",
                                                          "_selected_action": [quote.pk]})
        offer = Offer.objects.get(source_quote=quote)
        self.assertRedirects(r, f"/admin/offers/offer/{offer.pk}/change/", fetch_redirect_response=False)
        line = offer.lines.get()
        self.assertEqual((line.qty, line.unit_price), (D("50.000"), D("12.5000")))   # поточна ціна Minerva
        self.assertIn("у запиті: 11 EUR", line.price_note)

    def test_send_by_mail_marks_revision_sent(self):
        offer = services.new_offer(customer=self.customer, user=self.user, contact_name="Alex Buyer")
        services.add_line(offer, product=self.ant, qty=2)
        r = self.client.get("/email/compose/", {"offer": offer.pk})
        rev = OfferRevision.objects.get(offer=offer, revision=1)
        self.assertContains(r, f'name="att_ref" value="offer:{rev.pk}"')
        self.assertContains(r, 'value="buyer@client.example"')
        self.assertContains(r, f"Angebot {offer.number}")
        self.assertContains(r, "Sehr geehrte/r Alex Buyer")
        sent = {}

        def fake_send(client, **kw):
            sent.update(kw)
            return {"ok": True, "message_id": "<offer-mail@sevskiy.de>", "via": "test"}
        with patch("email_assistant.smtp_client.SMTPClient.send", fake_send):
            d = self.client.post("/email/send/", {"to": "buyer@client.example", "subject": f"Angebot {offer.number}",
                                                  "body": "anbei", "att_ref": [f"offer:{rev.pk}"]}).json()
        self.assertTrue(d["ok"], d)
        self.assertEqual(sent["attachments"][0]["name"], f"Angebot_{offer.number}.docx")
        rev.refresh_from_db()
        offer.refresh_from_db()
        self.assertEqual((rev.sent_message_id, offer.status), ("<offer-mail@sevskiy.de>", Offer.SENT))


class QuoteMailBase(OffersBase):
    """Скринька, клієнт, постачальник і помічник для листів — для тестів запитів пропозиції."""

    def setUp(self):
        super().setUp()
        from inventory.models import Supplier
        self.account = EmailAccount.objects.create(
            user=self.user, email_address=ME, imap_username=ME, imap_password="x",
            smtp_username=ME, smtp_password="x", imap_folder_inbox="INBOX", imap_folder_sent="Sent")
        es = EmailSettings.get_for_user(self.user)
        es.deadline_detection = es.telegram_notify_new = es.mark_read_on_server = False
        es.save()
        Supplier.objects.create(name="Parts Ltd", email="sales@parts.example")
        self._uid = 0

    def mail(self, subject, body, frm="buyer@client.example", folder="INBOX", to=(ME,), in_reply_to=""):
        self._uid += 1
        return save_message(self.account, {
            "uid": self._uid, "message_id": f"<qr{self._uid}@x.example>", "in_reply_to": in_reply_to,
            "references": [in_reply_to] if in_reply_to else [], "subject": subject, "from_email": frm,
            "from_name": "", "to_emails": list(to), "cc_emails": [], "body_text": body, "body_html": "",
            "attachments": [], "is_read": False, "sent_at": datetime(2026, 10, 1, 9, self._uid, tzinfo=dt_tz.utc)},
            folder)

    def qr(self, msg):
        from offers.models import QuoteRequest
        return QuoteRequest.objects.filter(message=msg).first()


class QuoteRequestTests(QuoteMailBase):
    """Запити пропозиції в пошті: розпізнавання, підказка в прев'ю, мітка й фільтр у списку, рішення людини."""

    def test_requests_in_several_languages(self):
        de = self.mail("Anfrage Antennen", "Guten Tag,\nbitte senden Sie uns Ihr Angebot für 500 Stück ANT-100.")
        en = self.mail("Inquiry", "Hello, could you please send us a quotation for CAB-7 x 20?", frm="ops@newco.example")
        uk = self.mail("Запит", "Просимо надіслати комерційну пропозицію на 50 шт ANT-100", frm="zakup@firma.example")
        price = self.mail("Frage", "Was kostet 200 Stk CAB-7?", frm="info@shop.example")
        for m in (de, en, uk, price):
            self.assertIsNotNone(self.qr(m), m.subject)
        self.assertEqual(self.qr(de).items, [{"product_id": self.ant.pk, "sku": "ANT-100", "qty": "500"}])
        self.assertIn("позиції: ANT-100 × 500", self.qr(de).reason)
        self.assertIn("артикули з кількістю й питання про ціну", self.qr(price).reason)

    def test_not_requests(self):
        thanks = self.mail("Re: Angebot", "Vielen Dank für Ihr Angebot, wir bestellen 500 Stück ANT-100.")
        supplier = self.mail("Unser Angebot", "Bitte beachten Sie unser Angebot für CAB-7.", frm="sales@parts.example")
        own = self.mail("Angebot A-2026-001", "Bitte finden Sie anbei unser Angebot.", frm=ME, folder="Sent",
                        to=("buyer@client.example",))
        news = save_message(self.account, {
            "uid": 900, "message_id": "<news@x.example>", "in_reply_to": "", "references": [], "subject": "Deals",
            "from_email": "news@shop.example", "from_name": "", "to_emails": [ME], "cc_emails": [],
            "body_text": "Request for quotation? Please quote us!", "body_html": "", "attachments": [],
            "is_read": False, "list_unsubscribe": True, "sent_at": datetime(2026, 10, 1, 8, 0, tzinfo=dt_tz.utc)},
            "INBOX")
        quoted = self.mail("Re: Anfrage", "Ok, passt so.\n\nAm 30.09.2026 schrieb Vertrieb:\n> Bitte Angebot für 500 Stück")
        chat = self.mail("Hallo", "Wie geht es Ihnen? Bis bald.")
        for m in (thanks, supplier, own, news, quoted, chat):
            self.assertIsNone(self.qr(m), m.subject)
        first = self.mail("Anfrage CAB", "Bitte Angebot für CAB-7 x 10")
        services.create_from_thread(first.thread, self.user)
        follow = self.mail("Re: Anfrage CAB", "Bitte auch ein Angebot für ANT-100", in_reply_to=first.message_id)
        self.assertEqual(follow.thread_id, first.thread_id)
        self.assertIsNone(self.qr(follow))                    # Angebot по гілці вже є — це переговори

    def test_preview_hint_list_tag_filter_and_dismiss(self):
        req = self.mail("Anfrage Antennen", "Bitte um ein Angebot für 500 Stück ANT-100")
        other = self.mail("Hallo", "Nur eine Info.")
        pv = self.client.get(f"/email/thread/{req.thread_id}/preview/")
        self.assertContains(pv, "Схоже на запит пропозиції")
        self.assertContains(pv, f"pvQuoteDismiss({self.qr(req).pk})")
        inbox = self.client.get("/email/")
        self.assertContains(inbox, "em-tag quote")
        only = self.client.get("/email/", {"quote": 1})
        self.assertContains(only, f'id="row-{req.pk}"')
        self.assertNotContains(only, f'id="row-{other.pk}"')
        r = self.client.post(f"/offers/api/quote-request/{self.qr(req).pk}/dismiss/")
        self.assertTrue(r.json()["ok"])
        self.assertEqual(self.qr(req).status, "dismissed")
        self.assertNotContains(self.client.get(f"/email/thread/{req.thread_id}/preview/"), "Схоже на запит пропозиції")
        self.assertNotContains(self.client.get("/email/", {"quote": 1}), f'id="row-{req.pk}"')
        from offers.detection import scan
        self.assertEqual(scan(req).status, "dismissed")       # рішення людини автоматика не перебиває
        stranger = User.objects.create_user("bob", "b@x.y", "p", is_staff=True)
        self.client.force_login(stranger)
        self.assertEqual(self.client.post(f"/offers/api/quote-request/{self.qr(req).pk}/dismiss/").status_code, 404)

    def test_angebot_from_thread_closes_request_and_backfill(self):
        from django.core.management import call_command
        from offers.models import QuoteRequest
        req = self.mail("Anfrage", "Please quote 500 pcs ANT-100")
        d = self.client.post(f"/offers/api/from-thread/{req.thread_id}/").json()
        self.assertEqual((self.qr(req).status, self.qr(req).offer_id), ("done", d["id"]))
        late = self.mail("RFQ", "RFQ: CAB-7 x 30", frm="ops@other.example")
        QuoteRequest.objects.filter(message=late).delete()     # як лист до появи розпізнавання
        out = io.StringIO()
        with patch("offers.detection.timezone.now", return_value=datetime(2026, 10, 15, tzinfo=dt_tz.utc)):
            call_command("detect_quote_requests", "--days", "30", stdout=out)
        self.assertIn("Знайдено запитів пропозиції: 1", out.getvalue())
        self.assertEqual(self.qr(late).status, "new")


class QuoteScanSettingsTests(QuoteMailBase):
    """«⚙️ Налаштування Angebote» → «Знайти запити в уже отриманій пошті»: глибина, фон, поле знову «—»."""

    def setUp(self):
        super().setUp()
        background = patch("offers.detection._start_background", lambda fn: fn())
        background.start()
        self.addCleanup(background.stop)

    def save_settings(self, depth):
        from offers.models import OfferSettings
        OfferSettings.get()
        with self.captureOnCommitCallbacks(execute=True):
            return self.client.post(f"/admin/offers/offersettings/1/change/", {
                "number_prefix": "A", "validity_days": 30, "default_currency": "EUR", "default_vat_rate": "19.00",
                "currency_by_country": "on", "currency_europe": "EUR", "currency_other": "USD",
                "price_mode": "auto", "lead_mode": "auto",
                "default_language": "de", "email_subject": "Angebot {number}", "quote_scan_depth": depth})

    def test_scan_old_mail_by_depth_and_switch_off(self):
        from offers.models import OfferSettings, QuoteRequest
        recent = self.mail("Anfrage", "Bitte Angebot für 500 Stück ANT-100")
        old = save_message(self.account, {
            "uid": 777, "message_id": "<old@x.example>", "in_reply_to": "", "references": [], "subject": "RFQ CAB-7",
            "from_email": "ops@old.example", "from_name": "", "to_emails": [ME], "cc_emails": [],
            "body_text": "Please quote 30 pcs CAB-7", "body_html": "", "attachments": [], "is_read": True,
            "sent_at": datetime(2024, 3, 1, 9, 0, tzinfo=dt_tz.utc)}, "INBOX")
        QuoteRequest.objects.all().delete()                     # як пошта, отримана до появи розпізнавання
        r = self.save_settings(30)
        self.assertEqual(r.status_code, 302, r.content[:400])
        s = OfferSettings.get()
        self.assertEqual(s.quote_scan_depth, OfferSettings.SCAN_OFF)          # одноразово — знову «—»
        self.assertTrue(s.quote_scan_status.startswith("✅"), s.quote_scan_status)
        self.assertIn("30 днів", s.quote_scan_status)
        self.assertIn("знайдено нових запитів 1", s.quote_scan_status)
        self.assertIsNotNone(self.qr(recent))
        self.assertIsNone(self.qr(old))                                        # старіший за 30 днів
        self.save_settings(OfferSettings.SCAN_ALL)
        self.assertIsNotNone(self.qr(old))
        self.assertIn("Уся пошта", OfferSettings.get().quote_scan_status)
        page = self.client.get("/admin/offers/offersettings/1/change/")
        self.assertContains(page, "Знайти запити в уже отриманій пошті")
        self.assertContains(page, '📝 Відкриті запити: <span id="qs-open">2</span>')

    def test_dismissed_requests_are_not_reopened(self):
        from offers.models import OfferSettings
        req = self.mail("Anfrage", "Bitte Angebot für 500 Stück ANT-100")
        self.client.post(f"/offers/api/quote-request/{self.qr(req).pk}/dismiss/")
        self.save_settings(OfferSettings.SCAN_ALL)
        self.assertEqual(self.qr(req).status, "dismissed")
        self.assertIn("знайдено нових запитів 0", OfferSettings.get().quote_scan_status)


class AdminTests(OffersBase):
    def test_add_with_inline_fills_catalog_price(self):
        from projects.models import Project
        project = Project.objects.create(name="Antennen", customer=self.customer)
        r = self.client.get("/admin/offers/offer/add/", {"project": project.pk})
        self.assertContains(r, "Client GmbH")                                  # клієнт із проєкту
        data = {
            "status": "draft", "customer": self.customer.pk, "contact_name": "Alex", "contact_email": "a@b.example",
            "billing_address": "x", "project": project.pk, "owner": self.user.pk, "language": "de",
            "currency": "EUR", "vat_rate": "19", "shipping_cost": "0", "valid_until": "2026-12-31",
            "price_mode": "auto", "lead_mode": "auto",
            "intro_text": "", "payment_terms": "", "delivery_terms": "", "outro_text": "", "internal_notes": "",
            "lines-TOTAL_FORMS": "1", "lines-INITIAL_FORMS": "0", "lines-MIN_NUM_FORMS": "0",
            "lines-MAX_NUM_FORMS": "1000",
            "lines-0-position": "1", "lines-0-product": self.ant.pk, "lines-0-sku": "", "lines-0-description": "",
            "lines-0-qty": "4", "lines-0-unit": "Stk", "lines-0-unit_price": "", "lines-0-discount_percent": "0",
            "lines-0-lead_time": "",
        }
        r = self.client.post("/admin/offers/offer/add/", data)
        self.assertEqual(r.status_code, 302, r.content[:800])
        offer = Offer.objects.get()
        line = offer.lines.get()
        self.assertEqual((line.sku, line.unit_price, line.price_source), ("ANT-100", D("12.5000"), "sale_price"))
        self.assertEqual(offer.gross_total, D("59.50"))
        page = self.client.get(f"/admin/offers/offer/{offer.pk}/change/")
        self.assertContains(page, "Сформувати документ")
        self.assertContains(page, "52.0 %")                                    # маржа (внутрішня)
        r = self.client.post(f"/admin/offers/offer/{offer.pk}/generate/")
        self.assertEqual(r.status_code, 302)
        self.assertTrue(OfferRevision.objects.filter(offer=offer, revision=1).exists())
        proj = self.client.get(f"/admin/projects/project/{project.pk}/change/")
        self.assertContains(proj, offer.number)


class ForwardedQuoteRequestTests(QuoteMailBase):
    """Колега переслав запит клієнта (FW / WG): відправник і текст — з пересланого листа."""

    FW = ("Mit freundlichen Grüßen\nSergey\n\n-------- Ursprüngliche Nachricht --------\n"
          "Von: Denys Repetynskyi <{frm}>\nDatum: 28.10.26 14:23\nAn: \"Sergey\" <" + ME + ">\nBetreff: Filters\n\n"
          "Good afternoon, we'd like to place a larger order.\n"
          "Please provide pricing and production lead times for 35 units and 100 units of ANT-100.\n")

    def test_forwarded_request_is_found(self):
        m = self.mail("FW: Filters", self.FW.format(frm="procurement10@moduletp.example"), frm=ME)
        self.assertEqual(m.category, "internal")
        qr = self.qr(m)
        self.assertIsNotNone(qr)
        self.assertTrue(qr.reason.startswith("переслано від procurement10@moduletp.example"))
        self.assertIn("Please provide pricing", qr.reason)

    def test_forwarded_from_supplier_or_colleague_is_not(self):
        supplier = self.mail("FW: Offer", self.FW.format(frm="sales@parts.example"), frm=ME)
        colleague = self.mail("FW: intern", self.FW.format(frm=ME), frm=ME)
        plain = self.mail("Notiz", "Please provide pricing — intern, ohne Weiterleitung", frm=ME)
        for m in (supplier, colleague, plain):
            self.assertIsNone(self.qr(m), m.subject)

    def test_colleague_forwards_own_forward(self):
        """«WG: FW: …» — колега переслав свій же FW: відправник — клієнт усередині, а не своя адреса."""
        inner = self.FW.format(frm="procurement10@moduletp.example")
        body = ("Mit freundlichen Grüßen\nDr. Sergey Sevskiy\n\n-----Ursprüngliche Nachricht-----\n"
                "Von: Dr. Sergey Sevskiy <" + ME + ">\nGesendet: Dienstag, 28. Oktober 2026 17:20\n"
                "An: slava@sevskiy.example\nBetreff: FW: Filters\n\n" + inner)
        m = self.mail("WG: FW: Filters", body, frm=ME)
        qr = self.qr(m)
        self.assertIsNotNone(qr)
        self.assertTrue(qr.reason.startswith("переслано від procurement10@moduletp.example"), qr.reason)
        data = services.quick_prefill(m.thread)
        self.assertEqual(data["contact_email"], "procurement10@moduletp.example")
        self.assertEqual(data["lines"][0]["qtys"], "35; 100")
        own_only = self.mail("WG: intern", body.replace("procurement10@moduletp.example", ME), frm=ME)
        from offers import detection
        self.assertIn("лише зі своїх адрес", detection.explain(own_only)["reason"])
        r = self.client.post(f"/offers/api/message/{own_only.pk}/rescan/").json()
        self.assertEqual((r["ok"], r["state"]), (True, "none"))

    def test_angebot_contact_and_customer_from_forwarded_sender(self):
        from crm.models import Customer
        buyer = Customer.objects.create(name="Denys", company="SL Tire", email="procurement10@moduletp.example")
        m = self.mail("FW: Filters", self.FW.format(frm="procurement10@moduletp.example"), frm=ME)
        offer = services.create_from_thread(m.thread, self.user)
        self.assertEqual((offer.contact_email, offer.contact_name, offer.customer),
                         ("procurement10@moduletp.example", "Denys", buyer))   # ім'я — з картки клієнта CRM


class QuickOfferTests(QuoteMailBase):
    """⚡ Легка форма Angebot: бланк Commercial Invoice, ступені (Staffelpreise), Deutsch + English, з листа."""

    def setUp(self):
        super().setUp()
        from shop.models import Shop, ShopListing, ShopPriceTier
        shop = Shop.objects.create(name="Website", slug="webshop", is_default=True, currency="EUR")
        listing = ShopListing.objects.create(shop=shop, product=self.ant, price=D("10.00"))
        ShopPriceTier.objects.create(listing=listing, min_qty=100, unit_price=D("8.00"))

    def test_parse_tiers(self):
        self.assertEqual(services.parse_tiers("35; 100"), [(D("35"), None), (D("100"), None)])
        self.assertEqual(services.parse_tiers("35, 100"), [(D("35"), None), (D("100"), None)])
        self.assertEqual(services.parse_tiers("35: 12,50; 100: 10,9"), [(D("35"), D("12.50")), (D("100"), D("10.9"))])
        self.assertEqual(services.parse_tiers(""), [])

    def test_tiers_price_each_quantity_and_document(self):
        from docx import Document
        offer = services.new_offer(customer=self.customer, user=self.user, language="de_en")
        offer = services.save_quick({"customer": self.customer, "language": "de_en"},
                                    [{"sku": "ANT-100", "qtys": "35; 100", "prices": ""},
                                     {"sku": "", "description": "Sonderanfertigung", "qtys": "1", "prices": "99,90"}],
                                    self.user, offer=offer)
        line = offer.lines.get(sku="ANT-100")
        self.assertEqual([(t["qty"], t["unit_price"]) for t in line.tiers], [("35", "10.0000"), ("100", "8.0000")])
        self.assertEqual((line.qty, line.unit_price, line.line_total), (D("35"), D("10.0000"), D("350.00")))
        offer.refresh_from_db()
        self.assertEqual(offer.net_total, D("449.90"))   # перша к-сть: 350 + 99,90
        rev = services.generate_document(offer)
        doc = Document(rev.docx_file.path)
        text = "\n".join(p.text for p in doc.paragraphs)
        cells = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn("Angebot / Quotation", text)                        # обома мовами
        self.assertIn("Menge, Einh.\nQty, Units", cells)
        rows = [[c.text for c in r.cells] for r in doc.tables[1].rows]
        self.assertIn(["1", "Antenna 2.4 GHz ANT-100", "35", "10,00", "350,00", ""], rows)   # кожна к-сть — позиція
        self.assertIn(["2", "Antenna 2.4 GHz ANT-100", "100", "8,00", "800,00", ""], rows)
        self.assertNotIn("Summe netto", cells)                            # ступені — варіанти, суми немає
        self.assertIn("Sevskiy GmbH", _all_text(rev.docx_file.path))     # бланк Commercial Invoice
        self.assertIn("Ship to (Receiver/Buyer):", cells)
        self.assertIn("Client GmbH", cells)
        self.assertNotIn("{{", text + cells)

    def test_prefill_from_forwarded_mail_and_create(self):
        from offers.models import QuoteRequest
        m = self.mail("FW: Filters", ForwardedQuoteRequestTests.FW.format(frm="procurement10@moduletp.example"), frm=ME)
        self.assertIsNotNone(self.qr(m))
        page = self.client.get(f"/offers/quick/?thread={m.thread_id}")
        self.assertEqual(page.status_code, 200)
        data = page.context["data"]
        self.assertEqual((data["language"], data["contact_email"]), ("en", "procurement10@moduletp.example"))
        self.assertEqual(data["lines"][0]["sku"], "ANT-100")
        self.assertEqual(data["lines"][0]["qtys"], "35; 100")             # «35 units and 100 units» → ступені
        import json
        r = self.client.post("/offers/quick/", {
            "thread": m.thread_id, "language": "en", "currency": "EUR", "valid_days": "30",
            "contact_name": "Denys", "contact_email": "procurement10@moduletp.example", "billing_address": "SL Tire",
            "lines_json": json.dumps(data["lines"])})
        offer = Offer.objects.latest("pk")
        self.assertRedirects(r, f"/offers/{offer.pk}/quick/", fetch_redirect_response=False)
        self.assertEqual((offer.language, offer.source_message_id, offer.lines.get().tiers[1]["qty"]),
                         ("en", m.pk, "100"))
        self.assertEqual(QuoteRequest.objects.get(message=m).status, QuoteRequest.DONE)
        self.assertTrue(offer.revisions.exists())                          # документ сформовано
        page = self.client.get(f"/offers/{offer.pk}/quick/")
        self.assertContains(page, "Перегенерувати")

    def test_price_hints_and_customer_api(self):
        d = self.client.get("/offers/api/price/?sku=ant-100&qtys=35;100&currency=EUR").json()
        self.assertEqual([t["unit_price"] for t in d["tiers"]], ["10", "8"])
        self.assertFalse(self.client.get("/offers/api/price/?sku=NOPE&qtys=1").json()["found"])
        c = self.client.get(f"/offers/api/customer/{self.customer.pk}/").json()
        self.assertIn("Hauptstr. 1", c["address"])

    def test_admin_tiers_field(self):
        offer = services.new_offer(customer=self.customer, user=self.user)
        line = services.add_line(offer, product=self.ant, qty=35)
        services.set_tiers(line, services.parse_tiers("35; 150"))
        line.save()
        self.assertEqual(services.tiers_text(line), "35: 10; 150: 8")


class ForwardedFormatsAndExplainTests(QuoteMailBase):
    """Формати пересилання різних програм, адреса з підпису, «чому не запит», хід пошуку наживо."""

    SIG = ("\n\nI look forward to your prompt response. Best regards,\n\n--\nBest regard,\n"
           "Denys Repetynskyi - supply manager Module ETP LLC\nModule ETP LLC\nPeremohy Ave 12\n03057 Kyiv, Ukraine\n"
           "Tel.: +380 44 123 45 67\nprocurement10@moduletp.example\n")

    def test_blank_lines_between_headers_thunderbird(self):
        body = ("Gruss\n\n-------- Weitergeleitete Nachricht --------\n\nBetreff: Filters\n\nDatum: Tue, 28 Oct 2026\n\n"
                "Von: Denys <procurement10@moduletp.example>\n\nAn: " + ME + "\n\n"
                "Please provide pricing for 35 units and 100 units of ANT-100." + self.SIG)
        m = self.mail("WG: Filters", body, frm=ME)
        self.assertIsNotNone(self.qr(m))

    def test_address_and_contact_from_signature(self):
        body = ForwardedQuoteRequestTests.FW.format(frm="procurement10@moduletp.example") + self.SIG
        m = self.mail("FW: Filters", body, frm=ME)
        data = services.quick_prefill(m.thread)
        self.assertEqual(data["contact_name"], "Denys Repetynskyi")
        self.assertEqual(data["billing_address"], "Module ETP LLC\nPeremohy Ave 12\n03057 Kyiv, Ukraine")

    def test_explain_why_not_and_rescan(self):
        from offers import detection
        from offers.models import QuoteRequest
        supplier = self.mail("FW: Offer", ForwardedQuoteRequestTests.FW.format(frm="sales@parts.example"), frm=ME)
        self.assertIn("постачальн", detection.explain(supplier)["reason"].lower())
        chat = self.mail("Hallo", "Wie geht es Ihnen?")
        self.assertEqual(detection.explain(chat)["state"], "none")
        first = self.mail("Anfrage CAB", "Bitte Angebot für CAB-7 x 10")
        services.create_from_thread(first.thread, self.user)
        follow = self.mail("Re: Anfrage CAB", "Bitte auch Angebot für ANT-100", in_reply_to=first.message_id)
        self.assertIn("вже є Angebot", detection.explain(follow)["reason"])
        req = self.mail("FW: Filters", ForwardedQuoteRequestTests.FW.format(frm="procurement10@moduletp.example"), frm=ME)
        QuoteRequest.objects.filter(message=req).delete()          # як лист, отриманий до оновлення
        self.assertEqual(detection.explain(req)["state"], "found")
        d = self.client.post(f"/offers/api/message/{req.pk}/rescan/").json()
        self.assertEqual(d["state"], "new")
        why = self.client.get(f"/email/message/{req.pk}/why/").json()
        self.assertEqual(why["quote"]["state"], "new")

    def test_scan_progress_live(self):
        from datetime import timedelta
        from django.utils import timezone
        from offers.models import OfferSettings
        for i in range(30):
            self.mail(f"Frage {i}", "Wie geht es?", frm=f"x{i}@client.example")
        steps = []
        from offers import detection
        detection.scan_recent(None, progress=lambda done, total, n: steps.append(done))
        self.assertIn(25, steps)                                     # кожні 25 листів, не лише наприкінці
        s = OfferSettings.get()
        OfferSettings.objects.filter(pk=s.pk).update(quote_scan_status="⏳ 30 днів: перевірено 25 з 343",
                                                     quote_scan_at=timezone.now())
        state = self.client.get("/admin/offers/offersettings/scan-status/").json()
        self.assertTrue(state["running"])
        page = self.client.get(f"/admin/offers/offersettings/{s.pk}/change/")
        self.assertContains(page, "scan-status/")
        OfferSettings.objects.filter(pk=s.pk).update(quote_scan_at=timezone.now() - timedelta(minutes=10))
        state = self.client.get("/admin/offers/offersettings/scan-status/").json()
        self.assertFalse(state["running"])
        self.assertIn("зупинилась", state["text"])

    def test_cascade_button_and_breaks(self):
        from shop.models import ShopSettings
        ShopSettings.get()
        page = self.client.get("/offers/quick/")
        self.assertContains(page, "ofCascade(")
        self.assertContains(page, 'id="of-breaks"')
        self.assertContains(page, "від 10 шт")


class OfferTemplateTests(QuoteMailBase):
    """📄 Word шаблон Angebot: поточний (бланк CI), приклад, свій з перевіркою, скидання; позиція з листа."""

    def test_download_sample_upload_reset(self):
        from io import BytesIO

        from django.core.files.uploadedfile import SimpleUploadedFile
        from docx import Document

        from offers.models import OfferSettings
        r = self.client.get("/offers/template/download/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("angebot_template.docx", r["Content-Disposition"])
        current = b"".join(r.streaming_content)
        self.assertIn("{{ L.ship_to }}", "".join(c.text for c in Document(BytesIO(current)).tables[0]._cells))   # бланк CI зі змінними
        sample = self.client.get("/offers/template/sample/")
        self.assertEqual(sample.status_code, 200)
        doc = Document(BytesIO(sample.content))
        rows = [[c.text for c in r.cells] for r in doc.tables[1].rows]
        self.assertIn(["1", "Filter FT220407", "35", "12,50", "437,50", "1-2"], rows)    # як на бланку Quotation
        self.assertIn(["2", "Filter FT220407", "100", "10,90", "1.090,00", "3-4"], rows)
        self.assertIn(["3", "Packaging and Shipping Costs", "1", "30,00", "30,00", ""], rows)
        bad = Document()
        bad.add_paragraph("{{ unclosed")
        buf = BytesIO()
        bad.save(buf)
        r = self.client.post("/offers/template/upload/", {"template": SimpleUploadedFile("x.docx", buf.getvalue())},
                             follow=True)
        self.assertContains(r, "Шаблон не прийнято")
        self.assertFalse(OfferSettings.get().template)
        r = self.client.post("/offers/template/upload/", {"template": SimpleUploadedFile("mine.docx", current)},
                             follow=True)
        self.assertContains(r, "Власний шаблон Angebot завантажено")
        self.assertTrue(OfferSettings.get().template)
        self.assertContains(self.client.get("/offers/quick/"), "власний (завантажений)")
        self.client.post("/offers/template/reset/")
        self.assertFalse(OfferSettings.get().template)
        page = self.client.get("/offers/quick/")
        self.assertContains(page, "на бланку Commercial Invoice")
        self.assertContains(page, "{{ l.qty }}")                       # довідник змінних

    def test_item_not_in_catalog_from_mail(self):
        body = ("Gruss\n\n-------- Ursprüngliche Nachricht --------\nVon: Denys <procurement10@moduletp.example>\n"
                "Betreff: Denys SL Tire\nGood afternoon, we'd like to place a larger order.\n\nFilter FT220407.\n"
                "Please provide pricing and production lead times for 35 units and 100 units.\n")
        m = self.mail("FW: Denys SL Tire", body, frm=ME)
        row = services.quick_prefill(m.thread)["lines"][0]
        self.assertEqual((row["sku"], row["description"], row["qtys"]), ("FT220407", "Filter FT220407", "35; 100"))


class PricingSourcesTests(QuoteMailBase):
    """Ціни Angebot: валюта за країною (Європа — EUR, інші — USD), склад Minerva → DigiKey, курс, строки зі складу,
    звідки що взято, 🏭 виробництво / розробка."""

    def setUp(self):
        super().setUp()
        from bots.models import DigiKeyConfig, DigiKeyListing
        from shop.models import ShopSettings
        ss = ShopSettings.get()
        ss.fx_rates = {"USD": "0.9"}                                   # 1 USD = 0,90 EUR
        ss.save()
        cfg = DigiKeyConfig.get()
        cfg.locale_currency = "USD"
        cfg.save()
        self.flt = Product.objects.create(sku="FLT-9", name="Filter 9", lead_time_days=21)   # без цін у Minerva
        DigiKeyListing.objects.create(product=self.flt, dk_title="Bandpass filter 9", dk_description="x",
                                      dk_prices=[{"qty": 1, "price": 20.0}, {"qty": 100, "price": 15.0}],
                                      dk_lead_time_days=11)
        DigiKeyListing.objects.create(dk_supplier_sku="DK-ONLY", dk_title="Only on DigiKey", dk_description="x",
                                      dk_prices=[{"qty": 10, "price": 5.0}])

    def test_currency_by_country(self):
        from offers.models import OfferSettings
        r = pricing.currency_for(self.customer)
        self.assertEqual(r["currency"], "EUR")
        self.assertIn("Німеччина", r["reason"])
        us = Customer.objects.create(name="Bob", company="US Inc", country="US")
        self.assertEqual(pricing.currency_for(us)["currency"], "USD")
        self.assertEqual(pricing.currency_for(None, address="Module ETP LLC\nKyiv 01001, Ukraine")["currency"], "EUR")
        self.assertEqual(pricing.currency_for(None, address="1 Main St\nAtlanta, Georgia 30301\nUSA")["currency"],
                         "USD")
        self.assertEqual(pricing.currency_for(None, email="a@shop.ca")["currency"], "USD")
        r = pricing.currency_for(None, email="a@moduletp.com")
        self.assertEqual(r["currency"], "EUR")                            # країну не визначено — з налаштувань
        self.assertIn("не визначено", r["reason"])
        self.assertEqual(services.defaults(us)["currency"], "USD")
        s = OfferSettings.get()
        s.currency_by_country = False
        s.save()
        self.assertEqual(pricing.currency_for(us)["currency"], "EUR")
        s.currency_by_country = True
        s.save()
        d = self.client.get(f"/offers/api/currency/?customer={us.pk}").json()
        self.assertEqual((d["currency"], d["country"]), ("USD", "US"))

    def test_stock_then_digikey_with_rate(self):
        q = pricing.quote_price(self.ant, 10, "USD")                      # склад Minerva: 12,50 EUR → USD
        self.assertEqual((q.unit_price, q.source), (D("13.8900"), OfferLine.SALE_PRICE))
        self.assertIn("× курс", q.note)
        self.assertIsNone(pricing.quote_price(self.ant, 10, "GBP").unit_price)   # курсу немає — без вигадок
        q = pricing.quote_price(self.flt, 150, "EUR")                     # у Minerva ціни немає → DigiKey
        self.assertEqual((q.unit_price, q.source), (D("13.5000"), OfferLine.DIGIKEY))
        self.assertIn("DigiKey Marketplace: від 100 шт", q.note)
        self.assertIsNone(pricing.quote_price(self.flt, 150, "EUR", mode="stock").unit_price)
        self.assertIsNone(pricing.quote_price(self.flt, 150, "EUR", mode="manual").unit_price)
        self.assertEqual(pricing.quote_price(self.ant, 1, "EUR", mode="digikey").source, OfferLine.SALE_PRICE)
        q = pricing.quote_price(None, 2, "USD", sku="dk-only")             # лише офер DigiKey, менше MOQ
        self.assertEqual(q.unit_price, D("5.0000"))
        self.assertIn("DigiKey продає від 10 шт", q.note)
        q = pricing.quote_price(None, 35, "EUR", sku="FT220407")
        self.assertIsNone(q.unit_price)
        self.assertIn("виробництво / розробку", q.note)

    def test_lead_time_from_stock_card_and_digikey_per_tier(self):
        from docx import Document

        from inventory.models import Location
        from inventory.services.stock import create_movement
        Location.objects.create(code="MAIN", location_type=Location.LocationType.FINISHED)
        create_movement(product=self.ant, tx_type="Incoming", qty=50)
        self.ant.lead_time_days = 14
        self.ant.save()
        offer = services.new_offer(customer=self.customer, user=self.user, language="de")
        offer = services.save_quick({"customer": self.customer, "language": "de"},
                                    [{"sku": "ANT-100", "qtys": "35; 100"},
                                     {"sku": "FLT-9", "qtys": "5", "lead_time": ""},
                                     {"sku": "CAB-7", "qtys": "10; 20", "lead_time": "1 Woche; 3 Wochen"}],
                                    self.user, offer=offer)
        ant = offer.lines.get(sku="ANT-100")
        self.assertEqual([t["lead_time"] for t in ant.tiers], ["ab Lager", "2-3"])   # 35 є, 100 — 14 днів ≈ 2-3 тижні
        self.assertEqual(ant.lead_time, "ab Lager")
        self.assertIn("доступно 50 шт", ant.lead_note)
        flt = offer.lines.get(sku="FLT-9")
        self.assertEqual((flt.lead_time, flt.price_source), ("3-4", OfferLine.DIGIKEY))   # картка товару, тижні
        self.assertEqual(flt.description, "Filter 9")
        cab = offer.lines.get(sku="CAB-7")
        self.assertEqual([t["lead_time"] for t in cab.tiers], ["1 Woche", "3 Wochen"])      # вручну через «;»
        rev = services.generate_document(offer)
        cells = "\n".join(c.text for t in Document(rev.docx_file.path).tables for r in t.rows for c in r.cells)
        rows = [[c.text for c in r.cells] for r in Document(rev.docx_file.path).tables[1].rows]
        self.assertEqual([r[5] for r in rows if r[1].startswith("Antenna")], ["ab Lager", "2-3"])
        self.assertEqual(services.quick_from_offer(offer)["lines"][2]["lead_time"], "1 Woche; 3 Wochen")
        self.assertEqual(services.quick_from_offer(offer)["lines"][0]["lead_time"], "")    # з Minerva — знову
        offer2 = services.new_offer(customer=self.customer, user=self.user, language="en", lead_mode="digikey")
        line = services.add_line(offer2, product=self.flt, qty=5)
        self.assertEqual(line.lead_time, "2-3")
        self.assertTrue(line.lead_note.startswith("DigiKey Marketplace"), line.lead_note)

    def test_production_request_kind(self):
        from offers.models import Offer
        offer = services.save_quick({"customer": self.customer, "language": "en"},
                                    [{"sku": "FT220407", "description": "Filter FT220407", "qtys": "35; 100"}],
                                    self.user)
        self.assertEqual(offer.request_kind, Offer.PRODUCTION)
        self.assertIn("виробництво", offer.lines.get().price_note)
        self.assertTrue(any("🏭" in w for w in services.warnings(offer)))
        offer = services.save_quick({"customer": self.customer, "language": "en"},
                                    [{"sku": "FT220407", "qtys": "35"}, {"sku": "ANT-100", "qtys": "1"},
                                     {"sku": "DK-ONLY", "qtys": "10"}], self.user, offer=offer)
        self.assertEqual(offer.request_kind, Offer.MIXED)
        self.assertEqual(offer.lines.get(sku="DK-ONLY").description, "Only on DigiKey")
        self.assertContains(self.client.get("/admin/offers/offer/?request_kind__exact=mixed"), offer.number)

    def test_quick_form_shows_where_data_came_from(self):
        body = ("Gruss\n\n-------- Ursprüngliche Nachricht --------\nVon: Denys <procurement10@moduletp.example>\n"
                "Betreff: Denys SL Tire\nGood afternoon, we'd like to place a larger order.\n\nFilter FT220407.\n"
                "Please provide pricing and production lead times for 35 units and 100 units.\n\n"
                "Best regards,\nDenys Repetynskyi - supply manager\nModule ETP LLC\nKyiv, Ukraine\n")
        m = self.mail("FW: Denys SL Tire", body, frm=ME)
        page = self.client.get(f"/offers/quick/?thread={m.thread_id}")
        data = page.context["data"]
        self.assertEqual(data["currency"], "EUR")
        self.assertIn("Україна", data["currency_reason"])
        self.assertEqual(data["origin"]["address"], "підпис листа")
        self.assertIn("кількості з листа: 35; 100", data["origin"]["lines"])
        self.assertContains(page, "↳ підпис листа")
        self.assertContains(page, "Строки поставки з")
        d = self.client.get("/offers/api/price/", {"sku": "FLT-9", "qtys": "1; 100", "currency": "EUR",
                                                    "language": "en"}).json()
        self.assertEqual([(t["unit_price"], t["group"]) for t in d["tiers"]], [("18", "digikey"), ("13.5", "digikey")])
        self.assertEqual(d["tiers"][0]["lead_time"], "3-4")
        self.assertEqual(d["digikey"]["title"], "Bandpass filter 9")
        d = self.client.get("/offers/api/price/", {"sku": "FT220407", "qtys": "35"}).json()
        self.assertEqual((d["found"], d["kind"]), (False, "production"))

    def test_all_quantities_for_each_item_and_codes_outside_catalog(self):
        body = ("-------- Original Message --------\nFrom: Denys <procurement10@moduletp.example>\nSubject: RFQ\n\n"
                "We'd like to place a larger order.\nFilter FT220407, antenna ANT-100 and filter FLT-9.\n"
                "Please provide pricing and production lead times for 35 units and 100 units.\n\n"
                "Best regards,\nDenys\nModule ETP LLC, VAT DE123456789\n")
        m = self.mail("FW: RFQ", body, frm=ME)
        rows = {r["sku"]: r for r in services.quick_prefill(m.thread)["lines"]}
        self.assertEqual(set(rows), {"ANT-100", "FLT-9", "FT220407"})          # VAT ID з підпису — не артикул
        self.assertEqual({r["qtys"] for r in rows.values()}, {"35; 100"})      # обидві кількості — кожній позиції
        self.assertEqual(rows["FT220407"]["description"], "Filter FT220407")

    def test_article_search_from_stock(self):
        Product.objects.create(sku="FT220407", name="Bandpass filter 2.2 GHz")
        Product.objects.create(sku="FT-2201", name="Filter 2.2", category="filter")
        Product.objects.create(sku="SOFT-1", name="Software licence")
        r = self.client.get("/offers/api/products/", {"q": "FT"}).json()["results"]
        self.assertEqual([x["sku"] for x in r[:2]], ["FT-2201", "FT220407"])        # артикул починається з FT
        self.assertIn("SOFT-1", [x["sku"] for x in r])                              # містить — нижче
        self.assertEqual(services.search_products("ft 2204")[0]["sku"], "FT220407")   # без розділювачів
        self.assertEqual([x["sku"] for x in services.search_products("filter 9")], ["FLT-9"])   # слова — у назві
        dk = services.search_products("only")
        self.assertEqual((dk[0]["sku"], dk[0]["source"]), ("DK-ONLY", "digikey"))     # офер DigiKey без товару
        self.assertEqual(services.search_products("  "), [])
