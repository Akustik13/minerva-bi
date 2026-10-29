"""Шаблон Angebot (DOCX з тегами docxtpl).

На бланку Commercial Invoice (`shipping.services.commercial_invoice_service.active_template_path` — стандартний
або завантажений власний) — як Quotation: увесь бланк лишається (колонтитули з логотипом і реквізитами, таблиця
адрес, заголовок, номер і дата, Billing address, таблиця позицій, Incoterms, підпис і печатка), тексти й теги CI
замінено на Angebot: «Quotation No. / Date / Valid until», таблиця Pos. | Description | Qty | Unit Value |
Total Value | Lead time (кожна кількість ступенів — окремий рядок, «Packaging and Shipping Costs» — рядок),
«All prices are exclusive of VAT…», Payment terms, «valid until …», Payment information, Name / Position.
Без бланка — простий шаблон (`python manage.py build_offer_template`). Змінні — `offers.document.build_context`.
"""
from copy import deepcopy

VERSION = "2"   # змінили побудову → кеш шаблону з бланка перебудовується

# текст у run-ах бланка CI → тег Angebot (форматування run-а лишається)
BLANK_RUNS = {
    "Ship to (Receiver/Buyer):": "{{ L.ship_to }}:",
    "{{ ship_to.company }}": "{{ customer_name }}",
    "{{ ship_to.street }}": "{{ billing_address }}",
    "{{ ship_to.city_line }}": "",
    "{{ ship_to.country }}": "",
    "Contact": "{{ L.contact_person }}",
    " Person": "",
    "{{ ship_to.contact }}": "{{ contact_name }}",
    "{{ ship_to.phone }}": "{{ contact_phone }}",
    "{{ ship_to.email }}": "{{ contact_email_line }}",
    "Ship from (Sender/Seller):": "{{ L.ship_from }}:",
}
# абзаци тіла бланка: текст run-а → тег Angebot (жирність і розмір run-а лишаються)
BODY_RUNS = {
    "Invoice": "{{ L.title }}",
    "Invoice No.: ": "{{ L.number }}: ", "{{ inv_number }}": "{{ number }}",
    "Invoice Date:": "{{ L.date }}:", "{{ inv_date }}": "{{ date }}",
    "Billing address: ": "{{ L.billing }}: ", "{{ billing_address }}": "{{ billing_line }}",
    "Terms of delivery (Incoterms):": "{{ L.incoterms }}:", "{{ incoterm_text }}": "{{ delivery_terms }}",
    "Name: {{ signer_name }}": "{{ L.name }}: {{ signer }}",
    "Position: {{ signer_position }}": "{{ L.position }}: {{ signer_position }}",
}
HEADER_RUNS = {"Invoice No.: {{ inv_number }}": "{{ L.number }}: {{ number }}"}


def _swap_runs(paragraphs, mapping):
    """Замінити текст run-ів за словником; абзаци, що стали порожніми через заміну, — прибрати."""
    for p in list(paragraphs):
        touched = False
        for r in p.runs:
            if r.text in mapping:
                r.text, touched = mapping[r.text], True
        if touched and not p.text.strip():
            p._p.getparent().remove(p._p)


def _swap_table(table, mapping):
    """Те саме у всіх комірках таблиці, включно з вкладеними (блок продавця на бланку — вкладена таблиця)."""
    for cell in table._cells:
        _swap_runs(cell.paragraphs, mapping)
        for inner in cell.tables:
            _swap_table(inner, mapping)


def _qn(tag):
    from docx.oxml.ns import qn
    return qn(tag)


def _set_par(p, text):
    """Текст абзацу одним run-ом (форматування першого run-а лишається)."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return p
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return p


def _clone(p, text, before=False):
    """Копія абзацу (той самий стиль) з іншим текстом — після p (або перед ним)."""
    from docx.text.paragraph import Paragraph
    el = deepcopy(p._p)
    for d in list(el.iter(_qn("w:drawing"))) + list(el.iter(_qn("w:pict"))):
        d.getparent().remove(d)
    (p._p.addprevious if before else p._p.addnext)(el)
    return _set_par(Paragraph(el, p._parent), text)


def _wrap(p, cond):
    """{%p if cond %} … {%p endif %} навколо абзацу."""
    _clone(p, "{%p if " + cond + " %}", before=True)
    _clone(p, "{%p endif %}")


def _has_drawing(p):
    return any(True for _ in p._p.iter(_qn("w:drawing"))) or any(True for _ in p._p.iter(_qn("w:pict")))


def _set_cell(cell, text):
    """Текст комірки одним run-ом першого абзацу (решта абзаців — геть)."""
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    _set_par(cell.paragraphs[0], text)


def _set_width(tc, width):
    tcpr = tc.get_or_add_tcPr()
    w = tcpr.find(_qn("w:tcW"))
    if w is None:
        w = tcpr.makeelement(_qn("w:tcW"), {})
        tcpr.append(w)
    w.set(_qn("w:w"), str(width))
    w.set(_qn("w:type"), "dxa")


def _items_table(table):
    """Таблиця позицій CI → Angebot: без «HS Code» і «Country of Origin», + «Lead time»; рядки — Angebot."""
    from docx.table import _Cell
    tbl = table._tbl
    header = [tc for tc in tbl.tr_lst[0].tc_lst]
    texts = ["".join(t.text or "" for t in tc.iter(_qn("w:t"))).lower() for tc in header]
    drop = [i for i, t in enumerate(texts) if "tariff" in t or "hs code" in t or "origin" in t]
    if any(len(tr.tc_lst) != len(header) for tr in tbl.tr_lst):
        drop = []                    # об'єднані комірки — колонки не чіпаємо
    cols = tbl.tblGrid.findall(_qn("w:gridCol"))
    widths = [int(c.get(_qn("w:w")) or 0) for c in cols]
    freed = sum(widths[i] for i in drop)
    for tr in tbl.tr_lst:
        tcs = tr.tc_lst
        for i in sorted(drop, reverse=True):
            tr.remove(tcs[i])
        tr.append(deepcopy(tr.tc_lst[-1]))          # «Lead time» — як остання колонка
    for i in sorted(drop, reverse=True):
        tbl.tblGrid.remove(cols[i])
    widths = [w for i, w in enumerate(widths) if i not in drop]
    lead = max(1100, min(1400, freed - 400)) if freed else 1100
    if len(widths) > 1:
        widths[1] = max(1500, widths[1] + freed - lead)     # «Description» ширша
    widths.append(lead)
    new_col = deepcopy(tbl.tblGrid.findall(_qn("w:gridCol"))[-1])
    tbl.tblGrid.append(new_col)
    for c, w in zip(tbl.tblGrid.findall(_qn("w:gridCol")), widths):
        c.set(_qn("w:w"), str(w))
    for tr in tbl.tr_lst:
        for tc, w in zip(tr.tc_lst, widths):
            _set_width(tc, w)
    rows = tbl.tr_lst

    def cells(tr):
        return [_Cell(tc, table) for tc in tr.tc_lst]

    def fill(tr, values):
        for cell, v in zip(cells(tr), values):
            if v is not None:
                _set_cell(cell, v)

    head = cells(rows[0])
    fill(rows[0], [None, "{{ H.description }}", "{{ H.qty }}", "{{ H.unit }}", "{{ H.total }}", "{{ H.lead }}"]
         if len(head) == 6 else [None] * len(head))
    if len(head) != 6:
        return
    marks = {i: "".join(t.text or "" for t in tr.iter(_qn("w:t"))) for i, tr in enumerate(rows)}
    loop = next((i for i, t in marks.items() if "{%tr for" in t), None)
    if loop is None:
        return
    fill(rows[loop], ["{%tr for l in rows %}", "", "", "", "", ""])
    fill(rows[loop + 1], ["{{ l.pos }}", "{{ l.description }}", "{{ l.qty }}", "{{ l.unit_price }}", "{{ l.total }}",
                          "{{ l.lead_time }}"])
    _plain_pos(cells(rows[loop + 1])[0])
    ship = next((i for i, t in marks.items() if "has_shipping" in t), None)
    if ship is not None and ship + 1 < len(rows):
        fill(rows[ship + 1], ["{{ shipping_pos }}", "{{ L.shipping_row }}", "1", "{{ shipping_cost }}",
                              "{{ shipping_cost }}", ""])
        pos_cell = cells(rows[ship + 1])[0]
        _plain_pos(pos_cell)
        size = next((r.font.size for p in cells(rows[loop + 1])[0].paragraphs for r in p.runs if r.font.size), None)
        for p in pos_cell.paragraphs:              # той самий розмір, що в номерах позицій
            for r in p.runs:
                r.font.size = size
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for tr in rows:                            # строк — по центру, як на бланку Quotation
        for p in cells(tr)[5].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    total = next((i for i, t in marks.items() if "total_amount" in t or "Total invoice" in t), None)
    if total is not None:
        tr = rows[total]
        fill(tr, ["", "{{ L.total_net }}", "", "", "{{ net_total }}", ""])
        endif = rows[ship + 2] if ship is not None and ship + 2 < len(rows) else rows[loop + 2]
        before, after = deepcopy(endif), deepcopy(endif)
        tr.addprevious(before)
        tr.addnext(after)
        _set_cell(_Cell(before.tc_lst[0], table), "{%tr if show_total %}")
        _set_cell(_Cell(after.tc_lst[0], table), "{%tr endif %}")


def _plain_pos(cell):
    """№ позиції — наш {{ l.pos }}, без автонумерації абзацу бланка (інакше Word показав би «11», «22»)."""
    for p in cell.paragraphs:
        ppr = p._p.get_or_add_pPr()
        for el in ppr.findall(_qn("w:numPr")) + ppr.findall(_qn("w:ind")) + ppr.findall(_qn("w:pStyle")):
            ppr.remove(el)                     # стиль списку бланка дає відступ ліворуч — номер виходить за комірку
        ind = ppr.makeelement(_qn("w:ind"), {_qn("w:left"): "0", _qn("w:right"): "0", _qn("w:firstLine"): "0"})
        jc = ppr.find(_qn("w:jc"))
        (jc.addprevious if jc is not None else ppr.append)(ind)


def _left(paragraphs):
    """Ліворуч, а не по ширині: рядки з розривами (адреса, реквізити) інакше розтягуються."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for p in paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _quotation_from_blank(blank):
    """Увесь бланк CI як Angebot (Quotation): тексти й теги CI → Angebot, порядок блоків — як на бланку."""
    from docx import Document
    from docx.text.paragraph import Paragraph
    doc = Document(str(blank))
    if len(doc.tables) < 2:
        return None
    _swap_table(doc.tables[0], BLANK_RUNS)
    for sec in doc.sections:
        for part in (sec.header, sec.first_page_header, sec.even_page_header):
            _swap_runs(part.paragraphs, HEADER_RUNS)
    items = doc.tables[1]
    paras = [Paragraph(el, doc._body) for el in doc.element.body if el.tag == _qn("w:p")]

    def find(prefix):
        return next((p for p in paras if p.text.strip().startswith(prefix)), None)
    date_p, billing_p, inco_p = find("Invoice Date"), find("Billing address"), find("Terms of delivery")
    reason_p, declare_p, pay_p = find("Reason of export"), find("I hereby declare"), find("Payment terms")
    from docx.shared import Cm, Pt
    if date_p is not None:                    # «Valid until» — під датою, тим самим форматуванням
        valid = deepcopy(date_p._p)
        date_p._p.addnext(valid)
        valid_p = Paragraph(valid, doc._body)
        _swap_runs([valid_p], {"Invoice Date:": "{{ L.valid_until }}:", "{{ inv_date }}": "{{ valid_until }}"})
        for p in (find("Invoice No."), date_p, valid_p):   # значення — однією колонкою (і для «Deutsch + English»)
            if p is not None:
                for pos in (4.2, 7.2):
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(pos))
    for cell in doc.tables[0]._cells:
        _left(cell.paragraphs)
        for inner in cell.tables:
            for c in inner._cells:
                _left(c.paragraphs)
    _swap_runs(paras, BODY_RUNS)
    if billing_p is not None:                 # проєкт — під адресою, лише якщо є
        proj = _clone(billing_p, "{{ L.project }}: {{ project }}")
        _wrap(proj, "project")
    first = items._tbl                         # вступ — над таблицею, лише якщо заповнено
    gap = Paragraph(deepcopy(paras[0]._p), doc._body)
    first.addprevious(gap._p)
    _set_par(gap, "{{ intro_text }}")
    for r in gap.runs:
        r.bold, r.font.size = False, None
    gap.alignment = None
    gap.paragraph_format.space_after = Pt(6)
    _wrap(gap, "intro_text")
    _items_table(items)
    after = items._tbl.getnext()
    if after is not None and after.tag == _qn("w:p"):   # ПДВ і брутто — під таблицею, лише без ступенів і з ПДВ
        vat = Paragraph(after, doc._body)
        line = _clone(vat, "{{ L.vat }} {{ vat_rate }} %: {{ vat_total }} {{ currency }}  ·  {{ L.gross }}: "
                           "{{ gross_total }} {{ currency }}", before=True)
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _wrap(line, "show_vat")
    if inco_p is not None:
        _wrap(inco_p, "delivery_terms")
    if reason_p is not None:
        _set_par(reason_p, "{{ L.vat_note }}")
    if declare_p is not None:
        _set_par(declare_p, "{{ L.payment }}: {{ payment_terms }}")
        _wrap(declare_p, "payment_terms")
    if pay_p is not None:                       # «Payment terms» бланка → «valid until», далі завершення й банк
        _set_par(pay_p, "{{ valid_text }}")
        outro = _clone(pay_p, "{{ outro }}")
        _wrap(outro, "outro")
        info = _clone(Paragraph(outro._p.getnext(), doc._body), "{%p if payment_info %}")   # після endif завершення
        blank1 = _clone(info, "")
        head = _clone(blank1, "{{ L.payment_info }}:")
        blank2 = _clone(head, "")
        bank = _clone(blank2, "{{ payment_info }}")
        for r in bank.runs:
            r.bold = True
        _left([pay_p, outro, head, bank])
        _clone(_clone(bank, "{%p endif %}"), "")
        # порожні абзаци між «Incoterms» і «valid until» — як на бланку Quotation, без пропусків
        start = reason_p if reason_p is not None else inco_p
        if start is not None:
            el = start._p.getnext()
            while el is not None and el is not pay_p._p:
                nxt = el.getnext()
                if el.tag == _qn("w:p"):
                    p = Paragraph(el, doc._body)
                    if not p.text.strip() and not _has_drawing(p):
                        el.getparent().remove(el)
                el = nxt
    pos = find("Position:") or next((p for p in paras if "{{ signer_position }}" in p.text), None)
    if pos is not None:
        _wrap(pos, "signer_position")
    return doc


def _from_blank(blank):
    """Бланк CI (VERSION 1): колонтитули й таблиця адрес лишаються, решта тіла документа прибирається."""
    from docx import Document
    doc = Document(str(blank))
    body = doc.element.body
    first = doc.tables[0]._tbl if doc.tables else None
    keep = True
    for el in list(body):
        if el.tag.endswith("}sectPr"):
            continue
        if first is not None and el is first:
            keep = False
            continue
        if not keep or first is None:
            body.remove(el)
    if first is not None:
        _swap_table(doc.tables[0], BLANK_RUNS)
    for sec in doc.sections:
        for part in (sec.header, sec.first_page_header, sec.even_page_header):
            _swap_runs(part.paragraphs, HEADER_RUNS)
    return doc


def build(path, blank=None):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    if blank:
        doc = _quotation_from_blank(blank)
        if doc is not None:               # увесь бланк перетворено — Angebot як Quotation
            doc.save(str(path))
            return
        doc = _from_blank(blank)          # незнайомий бланк — як раніше: шапка й таблиця адрес, далі свій текст
    else:
        doc = Document()
        sec = doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.8)
        style = doc.styles["Normal"]
        style.font.name, style.font.size = "Arial", Pt(9.5)

    def para(text="", *, size=None, bold=False, align=None, color=None, space_after=4):
        p = doc.add_paragraph()
        run = p.add_run(text)        # один run — теги docxtpl не розриваються
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if align:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        return p

    if blank:
        para("")
        para("{{ L.title }}", size=16, bold=True, space_after=2)
        para("{{ L.number }}: {{ number }}    {{ L.date }}: {{ date }}    {{ L.valid_until }}: {{ valid_until }}",
             color="333333")
        para("{%p if project %}")
        para("{{ L.project }}: {{ project }}", color="333333")
        para("{%p endif %}")
        para("")
    else:
        para("{{ company_block }}", size=8, color="555555", space_after=10)
        para("{{ L.title }} {{ number }}", size=16, bold=True, space_after=2)
        para("{{ L.date }}: {{ date }}    {{ L.valid_until }}: {{ valid_until }}", color="333333")
        para("{%p if project %}")
        para("{{ L.project }}: {{ project }}", color="333333")
        para("{%p endif %}")
        para("")
        para("{{ L.customer }}", bold=True, space_after=1)
        para("{{ customer_name }}", space_after=1)
        para("{{ billing_address }}", space_after=1)
        para("{%p if contact_name %}")
        para("{{ L.contact }}: {{ contact_name }}")
        para("{%p endif %}")
        para("")
    para("{{ L.greeting }}", space_after=2)
    para("{{ intro_text }}", space_after=8)

    headers = ["{{ L.pos }}", "{{ L.sku }}", "{{ L.description }}", "{{ L.qty }}", "{{ L.unit_price }}",
               "{{ L.discount }}", "{{ L.total }}"]
    widths = [Cm(1.0), Cm(3.2), Cm(5.8), Cm(1.8), Cm(2.3), Cm(1.5), Cm(2.4)]
    table = doc.add_table(rows=4, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cells_row = ["{{ l.pos }}", "{{ l.sku }}",
                 "{{ l.description }}{% if l.lead_time %} · {{ L.lead_time }}: {{ l.lead_time }}{% endif %}",
                 "{{ l.qty }}", "{{ l.unit_price }}", "{{ l.discount }}", "{{ l.total }}"]
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        c.width = widths[i]
    table.rows[1].cells[0].text = "{%tr for l in lines %}"
    for i, v in enumerate(cells_row):
        c = table.rows[2].cells[i]
        c.text = v
        c.width = widths[i]
        if i >= 3:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    table.rows[3].cells[0].text = "{%tr endfor %}"

    para("", space_after=6)
    right = WD_ALIGN_PARAGRAPH.RIGHT
    para("{%p if has_shipping %}")
    para("{{ L.shipping }}: {{ shipping_cost }} {{ currency }}", align=right, space_after=1)
    para("{%p endif %}")
    para("{{ L.net }}: {{ net_total }} {{ currency }}", align=right, space_after=1)
    para("{{ L.vat }} {{ vat_rate }} %: {{ vat_total }} {{ currency }}", align=right, space_after=1)
    para("{{ L.gross }}: {{ gross_total }} {{ currency }}", align=right, bold=True, size=11, space_after=10)
    para("{%p if payment_terms %}")
    para("{{ L.payment }}: {{ payment_terms }}", space_after=2)
    para("{%p endif %}")
    para("{%p if delivery_terms %}")
    para("{{ L.delivery }}: {{ delivery_terms }}", space_after=8)
    para("{%p endif %}")
    para("{{ outro_text }}", space_after=8)
    para("{{ L.closing }}", space_after=2)
    para("{%p if signer %}")
    para("{{ signer }}")
    para("{%p endif %}")

    if not blank:
        footer = doc.sections[0].footer.paragraphs[0]
        run = footer.add_run("{{ footer_text }}")
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string("666666")
    doc.save(str(path))
