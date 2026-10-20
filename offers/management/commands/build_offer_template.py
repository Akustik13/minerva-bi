"""Створити стандартний шаблон Angebot: offers/templates_docx/offer_template.docx (python-docx + теги docxtpl).

    python manage.py build_offer_template

Власний шаблон завантажується в «Налаштуваннях Angebote»; змінні — див. offers/document.py build_context.
"""
from django.core.management.base import BaseCommand

from offers.document import DEFAULT_TEMPLATE


def build(path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.8)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(9.5)

    def para(text="", *, size=None, bold=False, align=None, color=None, space_after=4):
        p = doc.add_paragraph()
        run = p.add_run(text)        # один run — теги docxtpl не розриваються
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if align:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        return p

    para("{{ company_block }}", size=8, color="555555", space_after=10)
    para("{{ L.title }} {{ number }}", size=16, bold=True, space_after=2)
    para("{{ L.date }}: {{ date }}    {{ L.valid_until }}: {{ valid_until }}", color="333333")
    para("{%p if project %}")
    para("{{ L.project }}: {{ project }}", color="333333")
    para("{%p endif %}")
    para("")
    para("{{ L.customer }}", bold=True, space_after=1)
    para("{{ customer_name }}", space_after=1)
    para("{{ billing_address }}", space_after=1)
    para("{%p if contact_name %}")
    para("{{ L.contact }}: {{ contact_name }}")
    para("{%p endif %}")
    para("")
    para("{{ L.greeting }}", space_after=2)
    para("{{ intro_text }}", space_after=8)

    headers = ["{{ L.pos }}", "{{ L.sku }}", "{{ L.description }}", "{{ L.qty }}", "{{ L.unit_price }}",
               "{{ L.discount }}", "{{ L.total }}"]
    widths = [Cm(1.0), Cm(3.2), Cm(5.8), Cm(1.8), Cm(2.3), Cm(1.5), Cm(2.4)]
    table = doc.add_table(rows=4, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cells_row = ["{{ l.pos }}", "{{ l.sku }}",
                 "{{ l.description }}{% if l.lead_time %} · {{ L.lead_time }}: {{ l.lead_time }}{% endif %}",
                 "{{ l.qty }}", "{{ l.unit_price }}", "{{ l.discount }}", "{{ l.total }}"]
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        c.width = widths[i]
    table.rows[1].cells[0].text = "{%tr for l in lines %}"
    for i, v in enumerate(cells_row):
        c = table.rows[2].cells[i]
        c.text = v
        c.width = widths[i]
        if i >= 3:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    table.rows[3].cells[0].text = "{%tr endfor %}"

    para("", space_after=6)
    right = WD_ALIGN_PARAGRAPH.RIGHT
    para("{%p if has_shipping %}")
    para("{{ L.shipping }}: {{ shipping_cost }} {{ currency }}", align=right, space_after=1)
    para("{%p endif %}")
    para("{{ L.net }}: {{ net_total }} {{ currency }}", align=right, space_after=1)
    para("{{ L.vat }} {{ vat_rate }} %: {{ vat_total }} {{ currency }}", align=right, space_after=1)
    para("{{ L.gross }}: {{ gross_total }} {{ currency }}", align=right, bold=True, size=11, space_after=10)
    para("{%p if payment_terms %}")
    para("{{ L.payment }}: {{ payment_terms }}", space_after=2)
    para("{%p endif %}")
    para("{%p if delivery_terms %}")
    para("{{ L.delivery }}: {{ delivery_terms }}", space_after=8)
    para("{%p endif %}")
    para("{{ outro_text }}", space_after=8)
    para("{{ L.closing }}")

    footer = sec.footer.paragraphs[0]
    run = footer.add_run("{{ footer_text }}")
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string("666666")
    doc.save(str(path))


class Command(BaseCommand):
    help = "Створити стандартний шаблон DOCX для Angebote"

    def handle(self, *args, **opts):
        DEFAULT_TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
        build(DEFAULT_TEMPLATE)
        self.stdout.write(f"Шаблон: {DEFAULT_TEMPLATE}")
