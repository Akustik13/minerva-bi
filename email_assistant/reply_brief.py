"""
«📋 Підготувати відповідь» — по кроках: дані → бриф → лист (ReplyBrief). До якого кроку йти, вибирає менеджер
(ReplyBrief.mode: лише дані / + бриф / + лист); далі — кнопками «📋 Скласти бриф» і «✍️ Написати лист».

1. gather() — БЕЗ AI (0 токенів), у момент натискання; кожне джерело отримує номер [#N] і групу походження:
   - mail       — листи цієї гілки до листа, на який відповідаємо;
   - history    — стан і хронологія, те саме, що в «📜 Історії» (замовлення, відправлення, нотатки CRM);
   - overview   — збережений AI-огляд проєкту (обіцянки, відкриті питання);
   - project    — шапка проєкту й останні листи з інших його гілок;
   - others     — інші листування з клієнтом з бази Minerva (останній лист кожного; це не RAG);
   - offer      — Angebote гілки / клієнта / проєкту з позиціями;
   - attachment — текст вкладень гілки (PDF, DOCX, TXT, CSV);
   - rag        — фрагменти з бази знань RAG, лише в листуваннях цього клієнта (або проєкту),
                  запит — побажання менеджера + суть листа.
   Без дублів (ReplyBrief.dedup — показується в «Зібраних даних»):
   - листи гілки, крім того, на який відповідаємо, — без процитованої історії (вона й так є окремими джерелами);
   - RAG не шукає в цій гілці й у гілках проєкту; листи, що вже є в джерелах, відкидаються;
     кілька фрагментів одного листа — одне джерело.
   Текст джерел зберігається в ReplyBrief.sources — бриф можна скласти пізніше без повторного збору.
2. Бриф — швидша модель (ai_helper.BRIEF_MODEL), українською, з посиланнями [#N].
3. Лист — модель із «Налаштувань пошти» за брифом; «↻ Переписати» / «✍️ Написати лист» — лише цей крок.
Ті самі зібрані дані — й для двох інших режимів у «📜 Історії» (замість брифу й листа):
   - answer — «❓ Запитати про лист»: відповідь на запитання менеджера з посиланнями [#N] (склад, історія,
     проблеми з товаром…); RAG — у всій пошті користувача, запит = запитання + суть листа + артикули;
   - review — «🧭 Що обіцяли / зробили»: обіцянки (хто, строк, статус), виконане, відкриті питання, наступна дія —
     той самий формат і перевірка, що огляд проєкту (projects.overview.validate), але для листування без проєкту.
Пересланий колегою лист (context.forwarded_origin) — RAG шукає за оригіналом у всій пошті, без фільтра клієнта.
Усе у фоні: браузер опитує стан (stages). Дані застарілі, якщо в гілці з'явились новіші листи.
"""
import json
import logging
import os
import re
import threading
import time

from django.conf import settings
from django.db import connection
from django.db.models import Q
from django.utils import timezone

logger = logging.getLogger('email_assistant')

THREAD_LAST = 12
NEWEST_CHARS, MESSAGE_CHARS = 10000, 2500
HISTORY_EVENTS = 40
OTHER_THREAD_CHARS = 800
PROJECTS, PROJECT_MESSAGES, PROJECT_CHARS = 2, 8, 1200
OFFERS, OFFER_LINES = 5, 12
ATTACHMENT_FILES, ATTACHMENT_CHARS, ATTACHMENT_TOTAL, ATTACHMENT_MAX_MB = 4, 4000, 12000, 15
ATTACHMENT_TYPES = ('pdf', 'docx', 'txt', 'csv')
RAG_HITS, RAG_HIT_CHARS, RAG_WAIT_SECONDS = 8, 1200, 12
RAG_REQUEST, RAG_LETTER_CHARS = 16, 2400   # фрагментів у запиті (із запасом на дублі) / тексту на один лист
QUOTE_NOTE = '[… цитату попередніх листів прибрано]'

STAGES = [
    ('thread', '📨 Ця гілка'),
    ('minerva', '📜 Історія (як у «📜 Історії»), огляд проєкту, інші листування з клієнтом'),
    ('projects', '📁 Проєкт: інші гілки'),
    ('offers', '📝 Angebote'),
    ('attachments', '📎 Вкладення'),
    ('rag', '🔎 RAG: пошук в інших листуваннях клієнта'),
    ('brief', '📋 Бриф (AI)'),
    ('letter', '✍️ Лист (AI)'),
    ('answer', '❓ Відповідь на запитання (AI)'),
    ('review', '🧭 Що обіцяли / зробили (AI)'),
]
GATHER_STAGES = [k for k, _ in STAGES if k not in ('brief', 'letter', 'answer', 'review')]
MODE_STEPS = {'data': ['brief', 'letter'], 'brief': ['brief', 'letter'], 'letter': ['brief', 'letter'],
              'answer': ['answer'], 'review': ['review']}
REPLY_MODES = ('data', 'brief', 'letter')


class Sources:
    """Нумеровані джерела [#N]: текст — для брифу й для показу менеджеру, група — звідки дані."""

    def __init__(self, items=None):
        self.items = list(items or [])

    def add(self, kind, label, text, url=''):
        self.items.append({'n': len(self.items) + 1, 'kind': kind, 'label': label[:200], 'url': url,
                           'text': (text or '').strip()})
        return len(self.items)

    def text(self) -> str:
        return '\n\n'.join(f'[#{i["n"]}] {i["label"]}\n{i.get("text") or "—"}' for i in self.items)

    def public(self, texts=True) -> list:
        keys = ('n', 'kind', 'label', 'url', 'text') if texts else ('n', 'kind', 'label', 'url')
        return [{k: i.get(k, '') for k in keys} for i in self.items]


def _dt(value) -> str:
    return timezone.localtime(value).strftime('%d.%m.%Y %H:%M') if value else '?'


def _num(value) -> str:
    if value is None:
        return '—'
    text = f'{value:f}'
    return text.rstrip('0').rstrip('.') if '.' in text else text


def _message_label(m, own) -> str:
    outgoing = m.folder == 'sent' or (m.from_email or '').lower() == own
    return f'{"→ НАДІСЛАЛИ МИ" if outgoing else "← ОТРИМАЛИ"} · {_dt(m.sent_at)} · {m.from_email} · «{m.subject or "—"}»'


def thread_messages(msg) -> list:
    """Листи гілки до листа, на який відповідаємо (як і для швидкої AI-відповіді)."""
    if not msg.thread_id:
        return [msg]
    msgs = list(msg.thread.messages.filter(is_deleted=False).order_by('sent_at', 'pk'))
    before = [m for m in msgs if not m.sent_at or not msg.sent_at or m.sent_at <= msg.sent_at] or [msg]
    return before[-THREAD_LAST:]


def latest_message_id(msg) -> int:
    if not msg.thread_id:
        return msg.pk
    return (msg.thread.messages.filter(is_deleted=False).order_by('-sent_at', '-pk')
            .values_list('pk', flat=True).first() or msg.pk)


# ── Текст вкладень ───────────────────────────────────────────────────────────

def _attachment_path(att) -> str:
    rel = att.get('file_path') or ''
    if not rel:
        return ''
    root = os.path.normpath(str(settings.MEDIA_ROOT))
    path = os.path.normpath(os.path.join(root, rel))
    return path if path.startswith(root + os.sep) and os.path.isfile(path) else ''


def attachment_text(att, limit=ATTACHMENT_CHARS, path=None) -> str:
    """Текст PDF / DOCX / TXT / CSV із збереженого вкладення; '' — якщо недоступне чи інший тип.
    path — уже перевірений шлях (файл із папки «📂 Джерела знань»), інакше — file_path у MEDIA."""
    name = (att.get('name') or '').lower()
    ext = name.rsplit('.', 1)[-1] if '.' in name else ''
    path = path if path is not None else _attachment_path(att)
    if ext not in ATTACHMENT_TYPES or not path or os.path.getsize(path) > ATTACHMENT_MAX_MB * 1024 * 1024:
        return ''
    try:
        if ext == 'pdf':
            from pypdf import PdfReader
            parts = []
            for page in PdfReader(path).pages[:15]:
                parts.append(page.extract_text() or '')
                if sum(len(p) for p in parts) > limit:
                    break
            text = '\n'.join(parts)
        elif ext == 'docx':
            import docx
            d = docx.Document(path)
            text = '\n'.join([p.text for p in d.paragraphs] +
                             [' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows])
        else:
            with open(path, 'rb') as fh:
                text = fh.read(limit * 4).decode('utf-8', errors='replace')
    except Exception as e:  # noqa: BLE001 — пошкоджений файл не зупиняє підготовку
        logger.info('reply brief: attachment %s not readable: %s', name, e)
        return ''
    return re.sub(r'[ \t]+', ' ', re.sub(r'\n{3,}', '\n\n', text)).strip()[:limit]


# ── Збір джерел (без AI) ─────────────────────────────────────────────────────

def gather(brief, on_stage=None) -> tuple:
    """(Sources, листи гілки, мітка листа-адресата). on_stage(key, status, detail) — для показу етапів."""
    from email_assistant.context import build_context
    on_stage = on_stage or (lambda *a: None)
    msg, user, wish = brief.message, brief.user, brief.wish
    account = msg.account
    own = (account.email_address or '').lower()
    src, used = Sources(), []
    dedup = {'quote_letters': 0, 'quote_chars': 0, 'rag_dropped': 0, 'rag_merged': 0}

    on_stage('thread', 'running')
    msgs = thread_messages(msg)
    target = ''
    for m in msgs:
        text = (m.body_text or '')[:NEWEST_CHARS] if m.pk == msg.pk else _without_quotes(m.body_text, dedup)
        n = src.add('mail', _message_label(m, own), text, f'/email/message/{m.pk}/')
        if m.pk == msg.pk:
            target = f'[#{n}] {_message_label(m, own)}'
    used.append(f'листів гілки: {len(msgs)}')
    on_stage('thread', 'done', f'{len(msgs)} лист.'
             + (f' · цитати прибрано в {dedup["quote_letters"]}' if dedup['quote_letters'] else ''))
    seen = {m.pk for m in msgs}                     # листи, що вже є в джерелах, — RAG їх не дублює

    on_stage('minerva', 'running')
    ctx = build_context(msg)
    seen |= {t['last'].pk for t in ctx.get('other_threads') or []}
    detail = _minerva(src, ctx)
    used += [d for d in detail if d]
    on_stage('minerva', 'done' if detail else 'skipped', ' · '.join(detail) or 'нічого пов\'язаного в базі')
    customer = ctx.get('customer')

    on_stage('projects', 'running')
    projects = _projects(msg)
    project_letters = 0
    for p in projects:
        head = (f'Проєкт {p.code} «{p.name}» · {p.get_status_display()}'
                + (f' · строк {p.due_date:%d.%m.%Y}' if p.due_date else '')
                + (f' · наступна дія: {p.next_action}' if p.next_action else ''))
        src.add('project', head, head, f'/admin/projects/project/{p.pk}/change/')
        letters = _project_letters(src, p, msg, own)
        seen |= letters
        project_letters += len(letters)
    if projects:
        used.append(f'проєкт {", ".join(p.code for p in projects)}: листів з інших гілок {project_letters}')
    on_stage('projects', 'done' if projects else 'skipped',
             f'{", ".join(p.code for p in projects)} · листів з інших гілок: {project_letters}' if projects
             else 'гілку не прив\'язано до проєкту')

    on_stage('offers', 'running')
    offers = _offers(src, msg, customer, projects)
    if offers:
        used.append(f'Angebote: {", ".join(offers)}')
    on_stage('offers', 'done' if offers else 'skipped', ', '.join(offers) or 'немає')

    on_stage('attachments', 'running')
    files = _attachments(src, msgs)
    if files:
        used.append(f'вкладення: {len(files)}')
    on_stage('attachments', 'done' if files else 'skipped', ', '.join(files) or 'немає текстових вкладень')

    on_stage('rag', 'running')
    status, detail, note = _rag(src, user, msg, ctx, wish, customer, projects, seen, dedup,
                                wide=brief.mode in ('answer', 'review'), everywhere=brief.mode == 'answer')
    if note:
        used.append(note)
    on_stage('rag', status, detail)

    brief.used, brief.dedup = used, dedup
    return src, msgs, target


def _without_quotes(body, dedup) -> str:
    """Лист гілки без процитованої історії: попередні листи гілки й так є окремими джерелами."""
    from email_assistant.quotes import new_text
    full = (body or '').strip()
    own = re.sub(r'[\s_]+$', '', new_text(full))
    if not own.strip() or len(own) >= len(full):    # цитати немає або лист — лише пересилання: лишаємо як є
        return full[:MESSAGE_CHARS]
    dedup['quote_letters'] += 1
    dedup['quote_chars'] += len(full) - len(own)
    return f'{own[:MESSAGE_CHARS]}\n{QUOTE_NOTE}'


def _minerva(src, ctx) -> list:
    """Дані з бази Minerva, розкладені за походженням (історія / огляд проєкту / інші листування)."""
    from email_assistant.context import _fmt, history_events, state_lines
    detail = []
    state, events = state_lines(ctx), history_events(ctx)
    if state or events:
        lines = (['ПОТОЧНИЙ СТАН:'] + state if state else []) + \
            (['ХРОНОЛОГІЯ:'] + [f'{_fmt(d)} {icon} {text}' + ('' if scope == 'order' else ' [інше по клієнту]')
                                for d, icon, text, scope in events[-HISTORY_EVENTS:]] if events else [])
        src.add('history', 'Історія: стан і хронологія (замовлення, відправлення, нотатки CRM)', '\n'.join(lines))
        if ctx.get('customer') is not None:
            detail.append(f'клієнт {ctx["customer"].name}')
        if ctx['orders']:
            detail.append(f'замовлень: {len(ctx["orders"])}')
        if ctx['notes']:
            detail.append(f'нотаток CRM: {len(ctx["notes"])}')
    if ctx.get('projects'):
        from projects.overview import facts_text
        for project, overview, stale in ctx['projects'][:PROJECTS]:
            when = timezone.localtime(overview.updated_at).strftime('%d.%m.%Y') if getattr(overview, 'updated_at', None) else ''
            src.add('overview', f'Огляд проєкту {project.code}' + (f' від {when}' if when else '')
                    + (' (застарів)' if stale else ''), facts_text(project, overview, stale)[:2500],
                    f'/admin/projects/project/{project.pk}/change/')
            detail.append(f'огляд {project.code}' + (' (застарів)' if stale else ''))
    for t in ctx.get('other_threads') or []:
        last = t['last']
        src.add('others', f'Інше листування «{t["thread"].subject or "—"}» ({t["count"]} лист.) — останній лист '
                          f'{_dt(last.sent_at)} від {last.from_email}', (last.body_text or '')[:OTHER_THREAD_CHARS],
                f'/email/message/{last.pk}/')
    if ctx.get('other_threads'):
        detail.append(f'інших листувань: {len(ctx["other_threads"])}')
    return detail


def _projects(msg) -> list:
    if not msg.thread_id:
        return []
    try:
        from projects.services import thread_projects
    except ImportError:
        return []
    found = {}
    for link in thread_projects(msg.thread)['confirmed']:
        found.setdefault(link.project_id, link.project)
    return list(found.values())[:PROJECTS]


def _project_letters(src, project, msg, own) -> set:
    """Останні листи інших гілок проєкту → джерела; повертає їхні pk."""
    from projects.services import project_messages
    qs = project_messages(project)
    if msg.thread_id:
        qs = qs.exclude(thread_id=msg.thread_id)
    letters = list(qs.order_by('-sent_at')[:PROJECT_MESSAGES])
    for m in reversed(letters):
        src.add('project', f'{project.code}: {_message_label(m, own)}', (m.body_text or '')[:PROJECT_CHARS],
                f'/email/message/{m.pk}/')
    return {m.pk for m in letters}


def _project_thread_ids(projects) -> list:
    """Підтверджені гілки проєктів — RAG у них не шукає: проєкт уже є окремим джерелом."""
    if not projects:
        return []
    from projects.models import ProjectLink
    return list(ProjectLink.objects.filter(project__in=projects, status=ProjectLink.CONFIRMED, thread__isnull=False)
                .values_list('thread_id', flat=True).distinct())


def _offers(src, msg, customer, projects) -> list:
    try:
        from offers.models import Offer
    except ImportError:
        return []
    cond = Q(source_message__thread_id=msg.thread_id) if msg.thread_id else Q(source_message=msg)
    if customer is not None:
        cond |= Q(customer=customer)
    if projects:
        cond |= Q(project__in=projects)
    numbers = []
    for o in Offer.objects.filter(cond).prefetch_related('lines').order_by('-created_at')[:OFFERS]:
        lines = []
        for line in list(o.lines.all())[:OFFER_LINES]:
            what = line.sku or (line.description or '').split('\n')[0][:60] or '—'
            lines.append(f'- {what}: {_num(line.qty)} {line.unit} × {_num(line.unit_price)} {o.currency}'
                         + (f', знижка {_num(line.discount_percent)}%' if line.discount_percent else '')
                         + (f', строк {line.lead_time}' if line.lead_time else '')
                         + (f' ({line.price_note})' if line.price_note else ''))
        text = (f'Статус: {o.get_status_display()} · створено {o.created_at:%d.%m.%Y}'
                + (f' · надіслано {o.sent_at:%d.%m.%Y}' if getattr(o, 'sent_at', None) else '')
                + (f' · дійсний до {o.valid_until:%d.%m.%Y}' if o.valid_until else '')
                + f' · нетто {_num(o.net_total)} {o.currency}, брутто {_num(o.gross_total)} {o.currency}\n'
                + '\n'.join(lines))
        src.add('offer', f'Angebot {o.number} (версія {o.revision})', text, f'/admin/offers/offer/{o.pk}/change/')
        numbers.append(o.number)
    return numbers


def _attachments(src, msgs) -> list:
    names, total = [], 0
    for m in reversed(msgs):                       # найновіші листи — першими
        for index, att in enumerate(m.attachments or []):
            if len(names) >= ATTACHMENT_FILES or total >= ATTACHMENT_TOTAL:
                return names
            text = attachment_text(att, min(ATTACHMENT_CHARS, ATTACHMENT_TOTAL - total))
            if text:
                src.add('attachment', f'Вкладення «{att.get("name")}» з листа {_dt(m.sent_at)} від {m.from_email}', text,
                        f'/email/message/{m.pk}/attachment/{index}/')
                names.append(att.get('name') or '?')
                total += len(text)
    return names


def _rag(src, user, msg, ctx, wish, customer, projects, seen=(), dedup=None, wide=False, everywhere=False) -> tuple:
    """RAG лише в листуваннях цього клієнта (або єдиного проєкту); → (status, detail для етапу, note для «врахувала»).
    У всій пошті користувача — для пересланого колегою листа (шукаємо за оригіналом) і коли менеджер сам питає
    (wide: «❓ Запитати», «🧭 Що обіцяли»). Без дублів: не шукає в цій гілці й гілках проєкту, відкидає листи
    з `seen` (уже в джерелах), фрагменти одного листа (чи одного вкладення) складає в одне джерело."""
    dedup = dedup if dedup is not None else {}
    try:
        from rag_assistant import retrieval
    except ImportError:
        return 'skipped', 'RAG не підключено', ''
    if not retrieval.available():
        return 'skipped', 'RAG не налаштований', ''
    if not retrieval.available_for(user):
        return 'skipped', 'немає права на пошук у базі знань', ''
    # запитання менеджера («які були проблеми з цим фільтром?») — про всіх клієнтів, тож без фільтра клієнта й проєкту
    customer_id = retrieval.customer_filter(ctx) if customer is not None and not everywhere else None
    project_id = (projects[0].pk if len(projects) == 1 and not customer_id and not ctx.get('customer_from_origin')
                  and not everywhere else None)
    wide = wide or bool(ctx.get('origin'))
    if not customer_id and not project_id and not wide:
        return 'skipped', 'клієнта чи проєкт не визначено — у всій пошті не шукаємо', ''
    query = retrieval.build_query(ctx, wish)   # стандартна частина з листа + побажання / запитання менеджера
    excluded = list(dict.fromkeys(([msg.thread_id] if msg.thread_id else []) + _project_thread_ids(projects)))[:100]
    scope = (f'клієнта {customer.name}' if customer_id else f'проєкту {projects[0].code}' if project_id
             else 'усій вашій пошті')
    return rag_sources(src, user, query, scope=scope, seen=seen, dedup=dedup, thread=msg.thread,
                       customer_id=customer_id, project_id=project_id, exclude_thread_ids=excluded)


def rag_sources(src, user, query, *, scope, seen=(), dedup=None, thread=None, customer_id=None, project_id=None,
                exclude_thread_ids=(), purpose='reply', seen_documents=()) -> tuple:
    """
    Пошук у RAG → джерела `rag` без дублів (спільне для «📋 Підготувати», огляду проєкту й карток проєкту / клієнта /
    замовлення / товару — rag_assistant/entity_help.py): листи з `seen`, документи з `seen_documents` («document:N»)
    і гілки exclude_thread_ids відкидаються, фрагменти одного листа (вкладення) чи документа — одне джерело.
    → (status, detail для етапу, note для «врахувала»).
    """
    from rag_assistant import retrieval
    dedup = dedup if dedup is not None else {}
    excluded = list(exclude_thread_ids)[:100]
    job = retrieval.start(user, query, purpose=purpose, thread=thread, customer_id=customer_id,
                          project_id=project_id, exclude_thread_ids=excluded, limit=RAG_REQUEST, wait_ms=8000)
    deadline = time.monotonic() + RAG_WAIT_SECONDS
    while not job.is_terminal and time.monotonic() < deadline:
        time.sleep(min(max(job.poll_after_ms or 1000, 300), 1500) / 1000)
        retrieval.poll(job)
    asked = f'запит «{query[:70]}{"…" if len(query) > 70 else ""}»'
    if job.status != job.COMPLETED:
        return 'skipped', f'{(job.message or "RAG не встиг")[:90]} · {asked}', ''
    letters, dropped, fragments, skip_threads = {}, 0, 0, set(excluded)
    for h in job.hits:
        att = h.get('attachment') or None
        document = h.get('kind') == 'document'
        if (h.get('message_pk') in seen and not att) or (h.get('thread_id') and h.get('thread_id') in skip_threads) \
                or h.get('document_id') in seen_documents:
            dropped += 1                                   # лист / документ уже є серед джерел (або виключена гілка)
            continue
        key = (h.get('url'), None) if document else (h.get('message_pk') or h.get('url'), att.get('index') if att else None)
        if key not in letters:
            if len(letters) >= RAG_HITS:
                continue
            letters[key] = (h, [])
        fragments += 1
        text = (h.get('text') or '').strip()[:RAG_HIT_CHARS]
        if text and text not in letters[key][1]:
            letters[key][1].append(text)
    for h, texts in letters.values():
        att = h.get('attachment') or None
        if h.get('kind') == 'document':     # файл / посилання бази знань: файл — сам документ, не вкладення
            label, att = (f'RAG · {h.get("document_kind") or "документ"} · {(h.get("sent_at") or "")[:10]} · '
                          f'«{(h.get("subject") or "")[:80]}»'), None
        else:
            label = f'RAG · {(h.get("sent_at") or "")[:10]} · {h.get("from_email", "")} · «{(h.get("subject") or "")[:80]}»'
        label = (label
                 + (f' · вкладення «{att.get("name") or "?"}»' if att else '')
                 + (f' · фрагментів {len(texts)}' if len(texts) > 1 else ''))
        src.add('rag', label, '\n[…]\n'.join(texts)[:RAG_LETTER_CHARS], h.get('url') or '')
    dedup['rag_dropped'] = dedup.get('rag_dropped', 0) + dropped
    dedup['rag_merged'] = dedup.get('rag_merged', 0) + fragments - len(letters)
    detail = (f'листів {len(letters)} (фрагментів {fragments}) у листуваннях {scope}'
              + (f' · відкинуто дублів {dropped}' if dropped else '') + f' · {asked}')
    return 'done', detail, f'RAG: листів {len(letters)}'


# ── Запуск ───────────────────────────────────────────────────────────────────

def _stages(keys=None, skip=()) -> list:
    return [{'key': k, 'label': label, 'status': 'skipped' if k in skip else 'pending',
             'detail': 'наступний крок — кнопкою нижче' if k in skip else ''}
            for k, label in STAGES if keys is None or k in keys]


def _restage(brief, run, skip=()) -> list:
    """Наступний крок: етапи збору лишаються (видно, звідки дані), кроки run — заново, skip — «кнопкою нижче»."""
    stages = brief.stages or _stages(GATHER_STAGES + MODE_STEPS['letter'])
    for s in stages:
        if s['key'] in run:
            s['status'], s['detail'] = 'pending', ''
        elif s['key'] in skip:
            s['status'], s['detail'] = 'skipped', 'наступний крок — кнопкою нижче'
    return stages


def _start_background(fn):
    threading.Thread(target=fn, daemon=True).start()


def _in_background(fn):
    caller = threading.current_thread()

    def run():
        try:
            fn()
        finally:
            if threading.current_thread() is not caller:   # власне з'єднання фонового потоку
                connection.close()
    _start_background(run)


def start(user, msg, *, wish='', recipient='', to_email='', cc=(), mode='letter'):
    """Нова підготовка для листа msg у фоні до кроку mode (data / brief / letter). → ReplyBrief (running)."""
    from email_assistant.models import ReplyBrief
    mode = mode if mode in dict(ReplyBrief.MODE_CHOICES) else ReplyBrief.MODE_LETTER
    skip = {'data': ('brief', 'letter'), 'brief': ('letter',)}.get(mode, ())
    keys = GATHER_STAGES + MODE_STEPS[mode]
    brief = ReplyBrief.objects.create(
        user=user, message=msg, thread=msg.thread if msg.thread_id else None, wish=(wish or '')[:2000], mode=mode,
        recipient=recipient[:300], to_email=to_email[:254], cc=list(cc), stages=_stages(keys, skip=skip),
        last_message_id=latest_message_id(msg))
    _in_background(lambda: _prepare(brief.pk))
    return brief


def start_brief(brief, wish=None):
    """«📋 Скласти бриф» зі збережених даних (без повторного збору). → той самий ReplyBrief (running)."""
    from email_assistant.models import ReplyBrief
    if wish is not None:
        brief.wish = wish[:2000]
    brief.status, brief.error, brief.stages = ReplyBrief.RUNNING, '', _restage(brief, ['brief'], skip=['letter'])
    brief.save()
    _in_background(lambda: _compose(brief.pk))
    return brief


def start_rewrite(brief, text, wish=None):
    """«↻ Переписати» / «✍️ Написати лист»: лише лист, з (виправленим) брифом. → той самий ReplyBrief (running)."""
    from email_assistant.models import ReplyBrief
    brief.brief = (text or '').strip()[:20000]
    if wish is not None:
        brief.wish = wish[:2000]
    brief.status, brief.error, brief.stages = ReplyBrief.RUNNING, '', _restage(brief, ['letter'])
    brief.save()
    _in_background(lambda: _rewrite(brief.pk))
    return brief


def _load(pk):
    from email_assistant.models import ReplyBrief
    return ReplyBrief.objects.select_related('message__account__user', 'message__thread', 'user').get(pk=pk)


def _stage_setter(brief):
    def set_stage(key, status, detail=''):
        for s in brief.stages:
            if s['key'] == key:
                s['status'], s['detail'] = status, (detail or '')[:220]
        brief.save(update_fields=['stages', 'updated_at'])
    return set_stage


def _model_label(model) -> str:
    from email_assistant.views import _model_label as label   # «Claude Opus 5.5» — як у «Minerva врахувала»
    return label(model) if model else '—'


def _prepare(pk):
    from email_assistant.models import ReplyBrief
    brief = _load(pk)
    set_stage = _stage_setter(brief)
    try:
        src, msgs, target = gather(brief, set_stage)
        brief.sources, brief.target = src.public(), target
        brief.save(update_fields=['sources', 'target', 'used', 'dedup', 'updated_at'])
        if brief.mode == ReplyBrief.MODE_ANSWER:
            _make_answer(brief, src, set_stage)
        elif brief.mode == ReplyBrief.MODE_REVIEW:
            _make_review(brief, src, set_stage)
        elif brief.mode != ReplyBrief.MODE_DATA:
            _make_brief(brief, src, set_stage)
            if brief.mode == ReplyBrief.MODE_LETTER:
                _write_letter(brief, msgs, set_stage)
        brief.status = ReplyBrief.DONE
    except Exception as e:  # noqa: BLE001
        logger.exception('reply brief %s failed: %s', pk, e)
        _fail(brief, e)
    brief.save()


def _compose(pk):
    from email_assistant.models import ReplyBrief
    brief = _load(pk)
    set_stage = _stage_setter(brief)
    try:
        _make_brief(brief, Sources(brief.sources), set_stage)
        brief.status = ReplyBrief.DONE
    except Exception as e:  # noqa: BLE001
        logger.exception('reply brief compose %s failed: %s', pk, e)
        _fail(brief, e)
    brief.save()


def _rewrite(pk):
    from email_assistant.models import ReplyBrief
    brief = _load(pk)
    set_stage = _stage_setter(brief)
    try:
        _write_letter(brief, thread_messages(brief.message), set_stage)
        brief.status = ReplyBrief.DONE
    except Exception as e:  # noqa: BLE001
        logger.exception('reply brief rewrite %s failed: %s', pk, e)
        _fail(brief, e)
    brief.save()


def _make_brief(brief, src, set_stage):
    from email_assistant import ai_helper
    set_stage('brief', 'running')
    profile = getattr(brief.user, 'profile', None)
    text = ai_helper.make_reply_brief(src.text(), brief.wish, brief.target, brief.message.account, profile)
    if not text:
        raise RuntimeError('AI не повернув бриф — перевірте ключ AI або спробуйте ще раз')
    brief.brief, brief.brief_model, brief.draft = text, ai_helper.last_model(), ''
    brief.used = [u for u in brief.used if not u.startswith(('бриф —', 'лист —'))] + \
        [f'бриф — {_model_label(brief.brief_model)}']
    set_stage('brief', 'done', f'модель {_model_label(brief.brief_model)}')


def _write_letter(brief, msgs, set_stage):
    from email_assistant import ai_helper
    from email_assistant.models import EmailMessage
    set_stage('letter', 'running')
    profile = getattr(brief.user, 'profile', None)
    draft = ai_helper.write_reply_from_brief(msgs, brief.message.account, profile, brief=brief.brief,
                                             wish=brief.wish, recipient=brief.recipient)
    if not draft:
        raise RuntimeError('AI не повернув лист — спробуйте «Переписати»')
    brief.draft, brief.letter_model = draft, ai_helper.last_model()
    brief.used = [u for u in brief.used if not u.startswith('лист —')] + \
        [f'лист — {_model_label(brief.letter_model)}']
    set_stage('letter', 'done', f'модель {_model_label(brief.letter_model)}')
    EmailMessage.objects.filter(pk=brief.message_id).update(ai_reply_draft=draft)


def _make_answer(brief, src, set_stage):
    """«❓ Запитати про лист»: відповідь на запитання менеджера лише за зібраними джерелами, з [#N]."""
    from email_assistant import ai_helper
    set_stage('answer', 'running')
    text = ai_helper.answer_from_sources(src.text(), brief.wish, brief.target, brief.message.account,
                                         getattr(brief.user, 'profile', None))
    if not text:
        raise RuntimeError('AI не відповів — перевірте ключ AI або спробуйте ще раз')
    brief.brief, brief.brief_model = text, ai_helper.last_model()
    brief.used = [u for u in brief.used if not u.startswith('відповідь —')] + \
        [f'відповідь — {_model_label(brief.brief_model)}']
    set_stage('answer', 'done', f'модель {_model_label(brief.brief_model)}')


def _make_review(brief, src, set_stage):
    """«🧭 Що обіцяли / зробили»: JSON як в огляді проєкту, перевірений; [#N] → лист чи документ із джерел."""
    from email_assistant import ai_helper
    from projects.overview import _parse_json, validate
    set_stage('review', 'running')
    raw = ai_helper.review_from_sources(src.text(), brief.target, brief.message.account,
                                        getattr(brief.user, 'profile', None))
    data = _parse_json(raw)
    if data is None:
        raise RuntimeError('AI не відповів у форматі огляду — спробуйте ще раз')
    by_n = {i['n']: i for i in src.items}

    def resolve(value):
        m = re.search(r'#\s*(\d+)', str(value or ''))
        item = by_n.get(int(m.group(1))) if m else None
        return {'label': f'#{item["n"]}', 'url': item.get('url') or ''} if item else {'label': '', 'url': ''}
    brief.brief = json.dumps(validate(data, [], {}, resolve=resolve), ensure_ascii=False)
    brief.brief_model = ai_helper.last_model()
    brief.used = [u for u in brief.used if not u.startswith('огляд —')] + [f'огляд — {_model_label(brief.brief_model)}']
    set_stage('review', 'done', f'модель {_model_label(brief.brief_model)}')


def review_display(brief):
    """Огляд «що обіцяли / зробили» для браузера: підписи статусів і дат, як у панелі огляду проєкту."""
    from projects.overview import for_display
    try:
        return for_display(json.loads(brief.brief)) if brief.brief else None
    except ValueError:
        return None


def _fail(brief, error):
    from email_assistant.models import ReplyBrief
    for s in brief.stages:
        if s['status'] == 'running':
            s['status'] = 'failed'
    brief.status, brief.error = ReplyBrief.ERROR, str(error)[:300]


# ── Для браузера ─────────────────────────────────────────────────────────────

def latest_for(user, msg, modes=REPLY_MODES):
    """Остання підготовка користувача для гілки листа (або для самого листа без гілки) серед режимів modes:
    відповідь листом (data / brief / letter) окремо від «❓ Запитати» (answer) і «🧭 Що обіцяли» (review)."""
    from email_assistant.models import ReplyBrief
    qs = ReplyBrief.objects.filter(user=user, mode__in=modes)
    qs = qs.filter(thread_id=msg.thread_id) if msg.thread_id else qs.filter(message=msg)
    return qs.select_related('message').first()


def as_json(brief) -> dict:
    """Стан для браузера; повний текст джерел — лише коли крок завершено (опитування під час роботи — легке)."""
    running = brief.status == brief.RUNNING
    stale = not running and latest_message_id(brief.message) != brief.last_message_id
    review = review_display(brief) if brief.mode == brief.MODE_REVIEW and not running else None
    return {'id': brief.pk, 'status': brief.status, 'mode': brief.mode, 'error': brief.error, 'stages': brief.stages,
            'brief': '' if review is not None else brief.brief, 'review': review,
            'sources': Sources(brief.sources).public(texts=not running), 'draft': brief.draft,
            'to': brief.to_email, 'cc': brief.cc, 'used': brief.used, 'dedup': brief.dedup or {}, 'wish': brief.wish,
            'message_id': brief.message_id, 'stale': stale,
            'updated': timezone.localtime(brief.updated_at).strftime('%d.%m %H:%M'),
            'gathered': timezone.localtime(brief.created_at).strftime('%d.%m.%Y %H:%M')}
