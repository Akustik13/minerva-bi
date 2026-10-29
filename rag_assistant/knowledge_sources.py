"""
«📂 Джерела знань» (етап 4): папки й сайти, які Minerva сама обходить і передає в базу знань (KnowledgeSource →
KnowledgeItem kind=source → тип document у стрічці для RAG; у Minerva — «❓ Запитати», огляди, картки).

Папка:
- лише всередині дозволених коренів (allowed_roots): <проєкт>/knowledge_sources (на NAS — у теці Minerva, у Docker уже
  видно), MEDIA/knowledge_sources і змонтовані шляхи з MINERVA_KNOWLEDGE_ROOTS; посилання (symlink) за межі — ні;
- файли НЕ копіюються: KnowledgeItem.source_path — шлях відносно папки, читається на місці (resolve);
- перечитується лише змінене (source_stamp = розмір + час зміни), зниклий файл → документ прибирається (tombstone);
- PDF / DOCX / TXT / CSV — текст витягує RAG (як вкладення листа), Excel / HTML / Markdown — Minerva сама (item.text);
- «📤 Завантажити файли / папку» з браузера — save_upload (назва без шляхів, тип і розмір джерела, вміст PDF / DOCX /
  XLSX за першими байтами, запис через тимчасовий файл), потім request_scan; папка без шляху — new_folder у
  MEDIA/knowledge_sources (managed: там же можна видалити файл — delete_file).
Сайт: посилання по одному в рядку (links_of); follow_links — sitemap.xml (у т. ч. вкладені) або сторінка з посиланнями
того самого сайту, інакше — лише ці сторінки / файли; до max_files; сторінку Minerva забирає сама (лише публічні адреси,
knowledge_items._download), HTML → текст, PDF → файл; сайт не відповів — його сторінки лишаються як були.
Прив'язки: клієнт / проєкт / товар з джерела для всіх документів + «артикул у назві файлу → товар».
Розклад: scan_due (cron `knowledge_sources --if-due`), «🔄 Перевірити зараз» — start_scan у фоні; одна перевірка
джерела одночасно (_claim — атомарний UPDATE), файли під час перевірки — rescan_requested → ще одна одразу після неї.
"""
import hashlib
import logging
import mimetypes
import os
import re
import threading
import uuid
import xml.etree.ElementTree as ET
from datetime import timedelta
from pathlib import Path
from urllib.parse import urljoin, urlparse

from django.conf import settings
from django.db import connection
from django.utils import timezone

logger = logging.getLogger(__name__)

DEFAULT_TYPES = ["pdf", "docx", "txt", "csv", "md", "html", "htm", "xlsx"]
MINERVA_TEXT = {"md", "markdown", "html", "htm", "xlsx"}     # текст витягує Minerva, а не RAG
TEXT_MAX = 100_000
SCAN_STALE = timedelta(hours=1)
BATCH = 300
_HREF_RE = re.compile(r"""href\s*=\s*["']([^"'#]+)""", re.I)
_SKU_TOKEN_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9._/-]{2,}")


class SourceError(Exception):
    """Зрозуміла людині причина (шлях поза коренями, папки немає, сайт не відповів…)."""


# ── Дозволені корені й шляхи ─────────────────────────────────────────────────

def allowed_roots() -> list:
    roots = [Path(settings.BASE_DIR) / "knowledge_sources", Path(settings.MEDIA_ROOT) / "knowledge_sources"]
    roots += [Path(p) for p in getattr(settings, "KNOWLEDGE_SOURCE_ROOTS", [])]
    out = []
    for r in roots:
        real = os.path.realpath(str(r))
        if real not in out:
            out.append(real)
    return out


def _inside(path, root) -> bool:
    try:
        return os.path.commonpath([path, root]) == root
    except ValueError:            # різні диски Windows
        return False


def check_folder(path) -> str:
    """Перевірена справжня адреса папки або SourceError з поясненням."""
    if not path or not os.path.isabs(path):
        raise SourceError("Потрібен повний шлях до папки, напр. /app/knowledge_sources/Dokumentation")
    real = os.path.realpath(path)
    roots = allowed_roots()
    if not any(_inside(real, r) for r in roots):
        raise SourceError("Папка поза дозволеними: " + ", ".join(roots)
                          + ". Інші папки NAS — змонтувати в Docker і додати в MINERVA_KNOWLEDGE_ROOTS.")
    if not os.path.isdir(real):
        raise SourceError("Папки немає (або Minerva її не бачить) — перевірте шлях і монтування в Docker.")
    return real


BROWSE_MAX = 300


def browse(path="") -> dict:
    """«📂 Вибрати на сервері»: дозволені корені (path порожній) або підпапки папки — лише в коренях, без прихованих
    і службових (@eaDir), посилання (symlink) не показуються. → {path, parent, folders [{name, path}], files}."""
    roots = [r for r in allowed_roots() if os.path.isdir(r)]
    if not (path or "").strip():
        return {"path": "", "parent": "", "folders": [{"name": r, "path": r} for r in roots], "files": 0}
    real = check_folder(path.strip())
    folders, files = [], 0
    try:
        with os.scandir(real) as entries:
            for e in sorted(entries, key=lambda e: e.name.lower()):
                if e.name.startswith((".", "@", "#", "~$")):
                    continue
                if e.is_dir(follow_symlinks=False):
                    if len(folders) < BROWSE_MAX:
                        folders.append({"name": e.name, "path": os.path.join(real, e.name)})
                elif e.is_file(follow_symlinks=False):
                    files += 1
    except OSError as e:
        raise SourceError(f"Папку не прочитати: {e.strerror or type(e).__name__}") from e
    parent = os.path.dirname(real)
    if real in roots or not any(_inside(parent, r) for r in roots):
        parent = ""
    return {"path": real, "parent": parent, "folders": folders, "files": files}


def resolve(source, rel) -> str:
    """Шлях файлу джерела — лише всередині його папки й дозволених коренів; '' — якщо файлу вже немає."""
    try:
        root = check_folder(source.path)
    except SourceError:
        return ""
    path = os.path.realpath(os.path.join(root, rel))
    return path if _inside(path, root) and os.path.isfile(path) else ""


def _types(source) -> set:
    return {t.strip().lower().lstrip(".") for t in (source.file_types or DEFAULT_TYPES) if str(t).strip()}


def _ext(name) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


# ── Завантаження з браузера ──────────────────────────────────────────────────

NAME_MAX = 180
DEPTH_MAX = 8
_BAD_CHARS_RE = re.compile(r'[\x00-\x1f<>:"/\\|?*]')
_MAGIC = {"pdf": b"%PDF", "docx": b"PK\x03\x04", "xlsx": b"PK\x03\x04"}


def safe_name(name) -> str:
    """Назва файлу чи папки без шляхів і заборонених символів (Windows / Synology); '' — прихований або службовий."""
    raw = str(name or "").strip()
    if not raw or raw.startswith((".", "~$", "@", "#")):
        return ""
    clean = " ".join(_BAD_CHARS_RE.sub(" ", raw).split()).rstrip(" .")
    if len(clean) > NAME_MAX:
        stem, dot, ext = clean.rpartition(".")
        clean = (stem[:NAME_MAX - len(ext) - 1] + "." + ext) if dot and len(ext) <= 10 else clean[:NAME_MAX]
    return clean


def _parts(path) -> list:
    return [p for p in re.split(r"[\\/]+", path or "") if p.strip()]


def _dirs(path) -> list:
    out = []
    for part in _parts(path):
        clean = safe_name(part)
        if not clean:
            raise SourceError(f"Папку «{part}» Minerva не бере (прихована або службова)")
        out.append(clean)
    return out


def new_folder(name) -> str:
    """Папка для джерела, куди файли завантажують у браузері: MEDIA/knowledge_sources/<назва> (завжди нова)."""
    base = Path(settings.MEDIA_ROOT) / "knowledge_sources"
    clean = safe_name(name) or "Dokumente"
    for n in range(1, 100):
        folder = base / (clean if n == 1 else f"{clean} ({n})")
        try:
            folder.mkdir(parents=True, exist_ok=False)
            return os.path.realpath(str(folder))
        except FileExistsError:
            continue
        except OSError as e:
            raise SourceError(f"Не вдалося створити папку: {e.strerror or type(e).__name__}") from e
    raise SourceError("Не вдалося створити папку — змініть назву")


def save_upload(source, upload, *, relpath="", subfolder="", replace=True) -> dict:
    """Файл із браузера → у папку джерела (relpath — шлях усередині вибраної папки). → {"path", "status":
    added | replaced | exists}; SourceError — чому ні (людською мовою)."""
    from .models import KnowledgeSource
    if source.kind != KnowledgeSource.FOLDER:
        raise SourceError("Файли завантажуються лише в джерело-папку")
    root = check_folder(source.path)
    name = safe_name((_parts(relpath) or [upload.name or ""])[-1])
    if not name:
        raise SourceError("Прихований або службовий файл — пропущено")
    types, ext = _types(source), _ext(name)
    if ext not in types:
        raise SourceError(f"Файли .{ext or '?'} у це джерело не беруться (типи: {', '.join(sorted(types))})")
    if not upload.size:
        raise SourceError("Порожній файл")
    if upload.size > source.max_file_mb * 1024 * 1024:
        raise SourceError(f"Файл більший за {source.max_file_mb} МБ")
    head = upload.read(8)
    upload.seek(0)
    if ext in _MAGIC and not head.startswith(_MAGIC[ext]):
        raise SourceError(f"Вміст файлу не схожий на .{ext}")
    dirs = _dirs(subfolder) + _dirs("/".join(_parts(relpath)[:-1]))
    if len(dirs) > DEPTH_MAX:
        raise SourceError(f"Задовгий шлях (понад {DEPTH_MAX} папок)")
    target = os.path.realpath(os.path.join(root, *dirs, name))
    folder = os.path.dirname(target)
    if not _inside(target, root):
        raise SourceError("Шлях поза папкою джерела")
    try:
        os.makedirs(folder, exist_ok=True)
    except OSError as e:
        raise SourceError(f"Не вдалося створити папку ({e.strerror or type(e).__name__}) — "
                          "папка лише для читання або немає прав") from e
    if not _inside(os.path.realpath(folder), root):
        raise SourceError("Шлях поза папкою джерела")
    exists = os.path.lexists(target)
    if exists and not os.path.isfile(target):
        raise SourceError("Там уже є папка з такою назвою")
    rel = "/".join(dirs + [name])
    if exists and not replace:
        return {"path": rel, "status": "exists"}
    tmp = os.path.join(folder, f".upload-{uuid.uuid4().hex}.part")     # «.» — перевірка папки його не бере
    try:
        with open(tmp, "wb") as fh:
            for chunk in upload.chunks():
                fh.write(chunk)
        os.replace(tmp, target)
    except OSError as e:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise SourceError(f"Не вдалося записати файл ({e.strerror or type(e).__name__}) — "
                          "папка лише для читання або немає прав") from e
    return {"path": rel, "status": "replaced" if exists else "added"}


def delete_file(source, rel) -> bool:
    """Видалити файл з папки, яку створила Minerva (лише там), і документ із бази знань (tombstone для RAG)."""
    from . import knowledge_items
    if not source.managed:
        raise SourceError("Видаляти файли можна лише в папці, яку створила Minerva; в інших — на NAS")
    root, path = check_folder(source.path), resolve(source, rel)
    if path:
        try:
            os.remove(path)
        except OSError as e:
            raise SourceError(f"Не вдалося видалити файл: {e.strerror or type(e).__name__}") from e
        folder = os.path.dirname(path)
        while folder != root and _inside(folder, root) and not os.listdir(folder):   # порожні підпапки
            os.rmdir(folder)
            folder = os.path.dirname(folder)
    for item in source.items.filter(source_path=rel):
        knowledge_items.delete(item)
    return bool(path)


# ── Текст, який Minerva витягує сама ─────────────────────────────────────────

def extract_text(path, ext) -> str:
    """Excel / HTML / Markdown → текст (PDF / DOCX / TXT / CSV читає RAG)."""
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        out, size = [], 0
        try:
            for ws in wb.worksheets:
                out.append(f"## {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() for c in row if c not in (None, "")]
                    if cells:
                        line = " | ".join(cells)
                        out.append(line)
                        size += len(line)
                    if size > TEXT_MAX:
                        break
                if size > TEXT_MAX:
                    break
        finally:
            wb.close()
        return "\n".join(out)[:TEXT_MAX]
    with open(path, "rb") as fh:
        raw = fh.read(TEXT_MAX * 4).decode("utf-8", errors="replace")
    if ext in ("html", "htm"):
        from email_assistant.imap_client import html_to_text
        raw = html_to_text(raw)
    return re.sub(r"\n{3,}", "\n\n", raw).strip()[:TEXT_MAX]


# ── Прив'язки ────────────────────────────────────────────────────────────────

class SkuIndex:
    """Артикули каталогу для «артикул у назві файлу → товар» (один запит на перевірку джерела)."""

    def __init__(self):
        from inventory.models import Product
        self.by_sku = {}
        for pk, sku, short in Product.objects.filter(is_active=True).values_list("pk", "sku", "sku_short"):
            for code in (sku, short):
                if code and len(code) >= 4:
                    self.by_sku.setdefault(code.lower(), pk)

    def find(self, text):
        for token in _SKU_TOKEN_RE.findall(text or ""):
            for t in (token.lower(), token.lower().strip("._-/")):
                if t in self.by_sku:
                    return self.by_sku[t]
            for part in re.split(r"[_\s/]+", token):         # «FT220407_Datenblatt», «Antennen/FT220407» → FT220407
                if part.lower() in self.by_sku:
                    return self.by_sku[part.lower()]
        return None


def _links(source, name, skus):
    """Прив'язки документа; артикул шукається спершу в назві файлу / сторінки, потім у шляху (назви папок)."""
    base = os.path.splitext(os.path.basename(urlparse(name).path if "://" in name else name))[0]   # «…FT220407.pdf»
    product_id = source.product_id or ((skus.find(base) or skus.find(name)) if skus else None)
    return {"customer_id": source.customer_id, "project_id": source.project_id, "product_id": product_id}


# ── Застосування ─────────────────────────────────────────────────────────────

def _apply(source, desired, fill) -> dict:
    """desired {source_path: spec}; fill(item, spec) → чи змінилось. Події стрічки — пачками в кінці."""
    from . import knowledge, knowledge_items
    from .models import KnowledgeItem
    existing = {i.source_path: i for i in source.items.all()}
    touched, created, updated = [], 0, 0
    with knowledge_items.quiet():
        for path, spec in desired.items():
            item = existing.pop(path, None)
            fresh = item is None
            if fresh:
                item = KnowledgeItem(kind=KnowledgeItem.SOURCE, auto=True, knowledge_source=source, source_path=path,
                                     owner_id=source.owner_id, source_key=f"src:{source.pk}:{path}"[:300])
            changed = fill(item, spec) or fresh
            if item.visibility != source.visibility or item.owner_id != source.owner_id:
                item.visibility, item.owner_id, changed = source.visibility, source.owner_id, True
            if changed:
                item.save()
                created += fresh
                updated += not fresh
                touched.append(item.pk)
        removed = list(existing.values())
        for item in removed:
            if item.file:                     # копія сторінки сайту (PDF) належить джерелу
                item.file.delete(save=False)
        KnowledgeItem.objects.filter(pk__in=[i.pk for i in removed]).delete()
    pks = touched + [i.pk for i in removed]
    for i in range(0, len(pks), BATCH):
        knowledge.touch_docs(knowledge_items.DOC_TYPE, pks[i:i + BATCH])
    return {"created": created, "updated": updated, "removed": len(removed)}


def _set(item, **values) -> bool:
    changed = False
    for field, value in values.items():
        if getattr(item, field) != value:
            setattr(item, field, value)
            changed = True
    return changed


def scan_folder(source) -> dict:
    root = check_folder(source.path)
    types, limit = _types(source), source.max_file_mb * 1024 * 1024
    skus = SkuIndex() if source.link_products_by_sku and not source.product_id else None
    files, skipped = {}, 0
    for dirpath, dirnames, filenames in os.walk(root, followlinks=False):
        dirnames[:] = sorted(d for d in dirnames if not d.startswith((".", "@", "#")))   # @eaDir Synology, .git…
        if not source.recursive:
            dirnames[:] = []
        for name in sorted(filenames):
            if name.startswith((".", "~$")) or _ext(name) not in types:
                continue
            full = os.path.join(dirpath, name)
            if not _inside(os.path.realpath(full), root):
                continue
            st = os.stat(full)
            if st.st_size > limit or st.st_size == 0:
                skipped += 1
                continue
            rel = os.path.relpath(full, root).replace(os.sep, "/")
            files[rel] = (full, st)
            if len(files) >= source.max_files:
                break
        if len(files) >= source.max_files:
            break

    def fill(item, spec):
        full, st = spec
        name = os.path.basename(full)
        stamp = f"{st.st_size}:{int(st.st_mtime)}"
        links = _links(source, item.source_path, skus)
        changed = _set(item, title=os.path.splitext(name)[0][:300] or name, file_name=name[:300], file_size=st.st_size,
                       content_type=(mimetypes.guess_type(name)[0] or "")[:200], **links)
        if item.source_stamp != stamp:
            item.source_stamp, item.file_sha256, changed = stamp, "", True
            ext = _ext(name)
            if ext in MINERVA_TEXT:
                try:
                    item.text = extract_text(full, ext)
                    item.fetch_error = ""
                except Exception as e:  # noqa: BLE001 — пошкоджений файл не зупиняє перевірку
                    item.text, item.fetch_error = "", f"Не вдалося прочитати: {type(e).__name__}"[:300]
        return changed

    result = _apply(source, files, fill)
    result["skipped"] = skipped
    return result


def links_of(source) -> list:
    """Посилання джерела-сайту: по одному в рядку, лише http(s), без повторів."""
    out = []
    for line in (source.url or "").splitlines():
        u = line.strip()
        if u and urlparse(u).scheme in ("http", "https") and urlparse(u).hostname and u not in out:
            out.append(u)
    return out


def _site_pages(start, limit) -> list:
    """Сторінки сайту від одного посилання: зі sitemap.xml (у т. ч. вкладених) або посилання сторінки того самого сайту."""
    from .knowledge_items import _download, _file_kind
    host = urlparse(start).hostname
    final, data, ctype = _download(start)
    urls = []
    if _file_kind(data, ctype):                     # посилання на PDF / DOCX — лише він
        return [final]
    if data.lstrip().startswith(b"<?xml") or "xml" in ctype:
        queue, seen = [data], set()
        while queue and len(urls) < limit:
            root = ET.fromstring(queue.pop(0))
            for loc in root.iter():
                if not loc.tag.endswith("loc") or not (loc.text or "").strip():
                    continue
                u = loc.text.strip()
                if u.endswith(".xml") and u not in seen and len(seen) < 20:     # індекс sitemap → вкладені
                    seen.add(u)
                    try:
                        queue.append(_download(u)[1])
                    except Exception:  # noqa: BLE001
                        continue
                elif urlparse(u).hostname == host:
                    urls.append(u)
    else:
        urls.append(final)
        for href in _HREF_RE.findall(data.decode("utf-8", errors="replace")):
            u = urljoin(final, href.strip())
            if urlparse(u).scheme in ("http", "https") and urlparse(u).hostname == host:
                urls.append(u.split("#")[0])
    return list(dict.fromkeys(urls))[:limit]


def _page_urls(source):
    """Сторінки всіх посилань джерела → (адреси, [(посилання, чому не відповіло)])."""
    from .knowledge_items import ItemError
    urls, failed = [], []
    for start in links_of(source):
        if len(urls) >= source.max_files:
            break
        if not source.follow_links:
            urls.append(start)
            continue
        try:
            urls += _site_pages(start, source.max_files - len(urls))
        except (ItemError, ET.ParseError) as e:
            failed.append((start, str(e) or type(e).__name__))
    return list(dict.fromkeys(urls))[:source.max_files], failed


def scan_website(source) -> dict:
    from . import knowledge_items
    skus = SkuIndex() if source.link_products_by_sku and not source.product_id else None
    urls, failed = _page_urls(source)
    if failed and not urls:
        raise SourceError(f"Не відповіло жодне посилання: {failed[0][1]} ({failed[0][0]})")
    pages = {u: u for u in urls}
    hosts = {urlparse(start).hostname for start, _ in failed}
    for path in source.items.values_list("source_path", flat=True):
        if path not in pages and urlparse(path).hostname in hosts:
            pages[path] = None                     # сайт не відповів — його сторінки лишаються як були

    def fill(item, url):
        if url is None:
            return False
        changed = _set(item, url=url[:2000], **_links(source, url, skus))
        previous = (item.text, item.file_sha256, item.title)
        knowledge_items.fetch(item, title_from_page=True)       # публічні адреси; HTML → текст, PDF → файл
        stamp = hashlib.sha256((item.text or item.file_sha256 or "").encode()).hexdigest()[:40]
        if item.source_stamp != stamp:
            item.source_stamp, changed = stamp, True
        return changed or previous != (item.text, item.file_sha256, item.title)

    result = _apply(source, pages, fill)
    if failed:
        result["failed"] = failed
    return result


# ── Запуск ───────────────────────────────────────────────────────────────────

RESCAN_ROUNDS = 3


def _claim(source) -> bool:
    """Атомарно позначити «йде перевірка»; False — уже йде (інший потік чи процес gunicorn / cron)."""
    from django.db.models import Q
    from .models import KnowledgeSource
    now = timezone.now()
    n = (KnowledgeSource.objects.filter(pk=source.pk)
         .filter(Q(scan_started_at__isnull=True) | Q(scan_started_at__lt=now - SCAN_STALE))
         .update(scan_started_at=now))
    if n:
        source.scan_started_at = now
    return bool(n)


def _scan_once(source) -> dict:
    from .models import KnowledgeSource
    started = timezone.now()
    try:
        result = scan_folder(source) if source.kind == KnowledgeSource.FOLDER else scan_website(source)
        failed = result.get("failed") or []
        text = (f"{timezone.localtime(started):%d.%m.%Y %H:%M}: документів {source.items.count()} · нових "
                f"{result['created']} · змінено {result['updated']} · прибрано {result['removed']}"
                + (f" · пропущено (розмір) {result['skipped']}" if result.get("skipped") else "")
                + (f" · не відповіли посилань {len(failed)}: {failed[0][0]} — {failed[0][1]}" if failed else ""))
        KnowledgeSource.objects.filter(pk=source.pk).update(last_scan_at=started, last_scan_result=text[:300],
                                                            last_error="")
        return result
    except Exception as e:  # noqa: BLE001
        message = str(e) if isinstance(e, (SourceError,)) else f"{type(e).__name__}: {e}"
        logger.warning("knowledge source %s scan failed: %s", source.pk, message)
        KnowledgeSource.objects.filter(pk=source.pk).update(last_scan_at=started, last_error=message[:300])
        return {"error": message}


def scan(source, claimed=False) -> dict:
    """Перевірити джерело зараз → лічильники; результат і помилка — у полях джерела. Файли, додані під час
    перевірки (rescan_requested), — ще одна перевірка одразу."""
    from .models import KnowledgeSource
    if not claimed and not _claim(source):
        return {"error": "Перевірка вже йде", "busy": True}
    try:
        for _ in range(RESCAN_ROUNDS):
            KnowledgeSource.objects.filter(pk=source.pk).update(rescan_requested=False)
            result = _scan_once(source)
            if not KnowledgeSource.objects.filter(pk=source.pk, rescan_requested=True).exists():
                break
            source.refresh_from_db()
        return result
    finally:
        KnowledgeSource.objects.filter(pk=source.pk).update(scan_started_at=None)


def running(source) -> bool:
    return bool(source.scan_started_at and timezone.now() - source.scan_started_at < SCAN_STALE)


def _start_background(fn):
    threading.Thread(target=fn, daemon=True).start()


def start_scan(source) -> bool:
    """«🔄 Перевірити зараз» у фоні; False — перевірка вже йде."""
    from .models import KnowledgeSource
    if not _claim(source):
        return False
    caller, pk = threading.current_thread(), source.pk

    def run():
        try:
            obj = KnowledgeSource.objects.filter(pk=pk).first()
            if obj is not None:
                scan(obj, claimed=True)
        finally:
            if threading.current_thread() is not caller:   # власне з'єднання фонового потоку
                connection.close()
    _start_background(run)
    return True


def request_scan(source) -> str:
    """Після завантаження файлів: перевірити зараз або одразу після поточної перевірки. → started | queued."""
    from .models import KnowledgeSource
    if start_scan(source):
        return "started"
    KnowledgeSource.objects.filter(pk=source.pk).update(rescan_requested=True)
    return "started" if start_scan(source) else "queued"     # перевірка щойно закінчилась — запустити самим


def scan_due() -> list:
    """Увімкнені джерела, які давно не перевіряли або чекають повторної перевірки (cron: knowledge_sources --if-due)."""
    from .models import KnowledgeSource
    done = []
    now = timezone.now()
    for source in KnowledgeSource.objects.filter(enabled=True).order_by("pk"):
        fresh = source.last_scan_at and now - source.last_scan_at < timedelta(hours=source.rescan_hours)
        if running(source) or (fresh and not source.rescan_requested):
            continue
        if not scan(source).get("busy"):
            done.append(source.pk)
    return done
